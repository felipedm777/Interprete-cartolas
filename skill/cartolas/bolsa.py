"""Dividendos y repartos que publica la Bolsa de Santiago.

La Bolsa publica cada día hábil un Excel, "resumen dividendos - repartos -
emisiones", con los dividendos y repartos de capital de las acciones, los
fondos de inversión y los emisores extranjeros que se transan en ella: por
evento, el nemotécnico, la moneda, el monto por acción o cuota, la fecha límite
y la fecha de pago. Es acumulado (trae desde 2022 hasta los pagos ya
anunciados), así que basta el más reciente.

Este módulo lo baja, lo lee y escribe la tabla de la Bolsa de la skill,
`datos/eventos_bolsa_<AAAAMMDD>.csv`, con la fecha del Excel en el nombre:

    python -m cartolas.bolsa                  # baja el más reciente
    python -m cartolas.bolsa <archivo.xlsx>   # usa uno ya bajado

Bajarlo necesita internet; leerlo, la librería openpyxl. La skill lee solo el
CSV. El Excel y la tabla tienen solo datos públicos de instrumentos.
"""

import argparse
import io
import re
import sys
import tempfile
import urllib.error
import urllib.request
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from . import eventos_capital as ec
from .numeros import leer_monto

URL = (
    "https://servicioscms.bolsadesantiago.com/Paginas/Descarga.aspx?attach="
    "Noticias%2Favisos+generales%2F{fecha:%Y%m%d}+resumen+dividendos+-+repartos+-+emisiones.xlsx"
)
CARPETA_DATOS = Path(__file__).parent / "datos"
_PREFIJO = "eventos_bolsa_"

# Hojas con eventos que pagan dinero, y la clase de evento de cada una. Las
# demás hojas (fechas sin derecho, emisiones de pago y liberadas, canjes,
# divisiones, stock splits, spin offs, llamados de capital, ofertas públicas y
# los dividendos de emisores colombianos del convenio Cross-Listing) no traen
# un monto en dinero por acción o cuota que se pueda cuadrar con un abono.
HOJAS = {
    "Dividendos acciones nacionales": ec.DIVIDENDO,
    "Dividendos CFI nacionales": ec.DIVIDENDO,
    "Repartos CFI-CFM": ec.REPARTO,
    "Dividendos emisores extranjeros": ec.DIVIDENDO,
}
# Cómo se nombra cada hoja en el texto del evento.
_CORTO = {
    "Dividendos acciones nacionales": "Acciones",
    "Dividendos CFI nacionales": "Dividendos CFI",
    "Repartos CFI-CFM": "Repartos CFI",
    "Dividendos emisores extranjeros": "Extranjeros",
}

# Columnas que se leen, por su encabezado. Las de nemotécnico son las que el
# encabezado deja entre "Serie" y "Tipo Instrumento"/"Mercado": la de pesos y
# la de dólares de los fondos, la del nemotécnico de las acciones (más una con
# el texto "Acciones") y la de dólares y pesos de los emisores extranjeros.
_EMISOR = "Nombre Emisor"
_TIPO = "Tipo (3)"
_MONEDA = "Moneda"
_MONTO = "Por Acción / cuota"
_LIMITE = ("Límite (1)", "Ex Date (1)")
_PAGO = "Pago"
_FIN_NEMOS = ("Tipo Instrumento", "Mercado")

_MONEDAS = {"$": "CLP", "US$": "USD", "EURO": "EUR", "SOLES": "PEN"}

# Montos escritos como texto que no son un número: la Bolsa remite al
# documento del emisor. Quedan sin monto, pero no son un dato mal leído.
_SIN_MONTO = {"", "-", "VER DOCUMENTO", "(VER OBSERVACIONES)"}


@dataclass
class Hoja:
    """Cuadratura de la lectura de una hoja."""

    nombre: str
    # Filas con contenido bajo el encabezado.
    filas: int = 0
    # Eventos leídos, y cuántos tienen monto por cuota.
    eventos: int = 0
    con_monto: int = 0
    # Filas que no son un evento (sin nemotécnico ni fecha de pago).
    descartadas: list[str] = field(default_factory=list)
    # Eventos con un valor que no se pudo leer con certeza: queda vacío y el
    # texto original va en el evento.
    dudosos: int = 0


@dataclass
class Lectura:
    al: date
    eventos: list[ec.Evento]
    hojas: list[Hoja]


def url(fecha: date) -> str:
    return URL.format(fecha=fecha)


def _es_xlsx(datos: bytes) -> bool:
    # Un xlsx es un zip. Para una fecha sin archivo la Bolsa responde 200 con
    # una página HTML.
    return datos[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(datos))


def descargar(carpeta: Path, hasta: date | None = None, dias: int = 10) -> Path:
    """Baja el Excel más reciente a `carpeta`, probando desde `hasta` hacia
    atrás, y devuelve su ruta (`<AAAAMMDD>.xlsx`)."""
    hasta = hasta or date.today()
    carpeta.mkdir(parents=True, exist_ok=True)
    fallas = []
    for i in range(dias + 1):
        fecha = hasta - timedelta(days=i)
        ruta = carpeta / f"{fecha:%Y%m%d}.xlsx"
        if ruta.exists() and _es_xlsx(ruta.read_bytes()):
            return ruta
        pedido = urllib.request.Request(url(fecha), headers={"User-Agent": "Mozilla/5.0"})
        try:
            with urllib.request.urlopen(pedido, timeout=60) as r:
                datos = r.read()
        except (urllib.error.URLError, OSError) as e:
            fallas.append(f"{fecha:%Y%m%d}: {e}")
            continue
        if _es_xlsx(datos):
            ruta.write_bytes(datos)
            return ruta
        fallas.append(f"{fecha:%Y%m%d}: no hay archivo")
    raise RuntimeError(
        f"no se encontró el Excel de la Bolsa entre el {hasta - timedelta(days=dias)} y el "
        f"{hasta}: " + "; ".join(fallas)
    )


def _fecha_excel(ruta: Path, libro) -> date:
    """Fecha del Excel: la del nombre que le da la Bolsa (o `descargar`), o la
    de su última modificación."""
    m = re.match(r"(\d{8})", ruta.name)
    if m:
        try:
            return datetime.strptime(m.group(1), "%Y%m%d").date()
        except ValueError:
            pass
    modificado = libro.properties.modified
    if modificado is None:
        raise ValueError(f"{ruta.name}: no se sabe la fecha del Excel; ponerla en el nombre (AAAAMMDD)")
    return modificado.date()


def _texto(valor) -> str:
    if valor is None:
        return ""
    # openpyxl deja los caracteres de control escapados ("CFIMO13B-E_x0006_").
    return re.sub(r"\s+", " ", re.sub(r"_x[0-9A-Fa-f]{4}_", "", str(valor))).strip()


def _fecha(valor) -> date | None:
    return valor.date() if isinstance(valor, datetime) else None


def _celda(fila: tuple, i: int):
    return fila[i] if i < len(fila) else None


def _monto(valor) -> tuple[Decimal | None, bool]:
    """Monto por cuota de la celda y si es dudoso.

    La celda es un número casi siempre. Cuando es texto, se acepta un número
    con coma decimal ("7,199693879412") o con punto decimal y más de tres
    decimales ("0.17699525"); otro texto ("2 18,1848371858") queda dudoso.
    """
    if isinstance(valor, bool):
        return None, True
    if isinstance(valor, (int, float)):
        return Decimal(repr(valor)) if isinstance(valor, float) else Decimal(valor), False
    t = _texto(valor).replace("\xa0", "")
    if t.upper() in _SIN_MONTO:
        return None, False
    if re.fullmatch(r"\d{1,3}(\.\d{3})*,\d+|\d+,\d+", t):
        return leer_monto(t), False
    if re.fullmatch(r"\d+\.\d{4,}", t):
        return Decimal(t), False
    return None, True


def _clase(hoja: str, tipo: str) -> str:
    """Clase del evento: la de la hoja, salvo un reparto informado entre los
    dividendos ("Reparto de Capital", "Reparto Final")."""
    t = tipo.upper()
    if HOJAS[hoja] == ec.DIVIDENDO and "REPARTO" in t:
        return ec.REPARTO
    if HOJAS[hoja] == ec.DIVIDENDO and "DISMINUC" in t:
        return ec.DISMINUCION
    return HOJAS[hoja]


def leer(ruta: Path) -> Lectura:
    """Lee los eventos con dinero del Excel de la Bolsa."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("leer el Excel de la Bolsa necesita openpyxl: pip install openpyxl") from None
    import warnings

    with warnings.catch_warnings():
        # El Excel trae áreas de impresión que openpyxl no entiende.
        warnings.simplefilter("ignore")
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        al = _fecha_excel(ruta, libro)
        faltan = [h for h in HOJAS if h not in libro.sheetnames]
        if faltan:
            raise ValueError(f"{ruta.name}: faltan las hojas {', '.join(faltan)}")
        eventos: list[ec.Evento] = []
        hojas = []
        for nombre in HOJAS:
            hoja, leidos = _leer_hoja(nombre, libro[nombre].iter_rows(values_only=True))
            hojas.append(hoja)
            eventos += leidos
    finally:
        libro.close()
    return Lectura(al, eventos, hojas)


def _leer_hoja(nombre: str, filas) -> tuple[Hoja, list[ec.Evento]]:
    hoja = Hoja(nombre)
    filas = iter(filas)
    for fila in filas:
        if fila and _texto(fila[0]) == _EMISOR:
            encabezado = [_texto(v) for v in fila]
            break
    else:
        raise ValueError(f"hoja {nombre}: no se encontró el encabezado")
    col = {}
    for i, c in enumerate(encabezado):
        col.setdefault(c, i)
    fin = next((col[c] for c in _FIN_NEMOS if c in col), None)
    limite = next((col[c] for c in _LIMITE if c in col), None)
    faltan = [c for c in (_TIPO, _MONEDA, _MONTO, _PAGO) if c not in col]
    if fin is None or limite is None or "Serie" not in col or faltan:
        raise ValueError(f"hoja {nombre}: el encabezado no tiene las columnas esperadas")
    columnas_nemo = range(col["Serie"] + 1, fin)

    eventos: list[ec.Evento] = []
    vistos = set()
    for fila in filas:
        if not any(_texto(v) for v in fila):
            continue
        hoja.filas += 1
        nemos = []
        for i in columnas_nemo:
            # Una celda puede traer dos nemotécnicos ("WMT; WALMARTCL").
            for n in re.split(r"[;/,]", _texto(_celda(fila, i)).upper()):
                n = re.sub(r"\s+", "", n)
                # La columna extra de las acciones dice "Acciones"; las demás
                # pueden traer "-" o un número suelto.
                if (
                    re.fullmatch(r"[A-Z0-9][A-Z0-9.\-]*", n)
                    and n != "ACCIONES"
                    and not n.isdigit()
                    and ec._nemo_clave(n) not in {ec._nemo_clave(x) for x in nemos}
                ):
                    nemos.append(n)
        pago = _fecha(_celda(fila, col[_PAGO]))
        if not nemos or pago is None:
            hoja.descartadas.append(
                " | ".join(_texto(v)[:40] for v in fila if _texto(v)) or "(vacía)"
            )
            continue
        tipo = _texto(_celda(fila, col[_TIPO]))
        moneda_texto = _texto(_celda(fila, col[_MONEDA]))
        moneda = _MONEDAS.get(moneda_texto.upper(), "")
        monto, dudoso = _monto(_celda(fila, col[_MONTO]))
        monto_texto = _texto(_celda(fila, col[_MONTO]))
        texto = f"{_CORTO[nombre]}: {tipo or '-'}"
        if monto is None and monto_texto:
            texto += f" (monto dudoso: {monto_texto})" if dudoso else f" ({monto_texto})"
        if moneda_texto not in ("", "-") and not moneda:
            dudoso = True
            texto += f" (moneda dudosa: {moneda_texto})"
        hoja.eventos += 1
        hoja.con_monto += monto is not None
        hoja.dudosos += dudoso
        for n in nemos:
            evento = ec.Evento(
                rut_fondo="",
                nemotecnico=n,
                clase=_clase(nombre, tipo),
                moneda=moneda,
                monto_por_cuota=monto,
                fecha_limite=_fecha(_celda(fila, limite)),
                fecha_pago=pago,
                texto=texto,
                fuente=ec.BOLSA,
            )
            clave = (ec._nemo_clave(n), evento.clase, moneda, monto, evento.fecha_limite, pago)
            if clave not in vistos:
                vistos.add(clave)
                eventos.append(evento)
    return hoja, eventos


# --- La tabla de la Bolsa en la skill -----------------------------------------


def ruta_tabla(al: date, carpeta: Path = CARPETA_DATOS) -> Path:
    return carpeta / f"{_PREFIJO}{al:%Y%m%d}.csv"


def tabla_vigente(carpeta: Path = CARPETA_DATOS) -> Path | None:
    """La tabla de la Bolsa más reciente de la carpeta."""
    tablas = sorted(carpeta.glob(f"{_PREFIJO}[0-9]*.csv"))
    return tablas[-1] if tablas else None


def _fecha_tabla(ruta: Path) -> date | None:
    m = re.search(r"(\d{8})", ruta.stem)
    return datetime.strptime(m.group(1), "%Y%m%d").date() if m else None


def cargar(ruta: Path | None = None) -> ec.Tabla:
    """Tabla de la Bolsa: la indicada (CSV de la skill o Excel de la Bolsa) o
    la más reciente de la skill. Vacía si no hay ninguna."""
    ruta = ruta or tabla_vigente()
    if ruta is None or not ruta.exists():
        return ec.Tabla([])
    if ruta.suffix.lower() == ".xlsx":
        lectura = leer(ruta)
        return ec.Tabla(lectura.eventos, lectura.al)
    return ec.Tabla.cargar(ruta, fuente=ec.BOLSA, bolsa_al=_fecha_tabla(ruta))


def escribir(lectura: Lectura, desde: date | None, carpeta: Path = CARPETA_DATOS) -> Path:
    """Escribe la tabla de la Bolsa con los eventos pagados desde `desde` y
    borra las tablas anteriores de la carpeta."""
    eventos = [e for e in lectura.eventos if desde is None or e.fecha_pago >= desde]
    ruta = ruta_tabla(lectura.al, carpeta)
    ec.escribir(eventos, ruta)
    for anterior in carpeta.glob(f"{_PREFIJO}[0-9]*.csv"):
        if anterior != ruta:
            anterior.unlink()
    return ruta


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m cartolas.bolsa",
        description="Baja el Excel de dividendos y repartos de la Bolsa de Santiago (o lee uno ya "
        "bajado) y escribe la tabla de la Bolsa de la skill.",
    )
    parser.add_argument("excel", nargs="?", help="Excel ya bajado; si no se indica, se baja el más reciente")
    parser.add_argument(
        "--carpeta",
        default=str(Path(tempfile.gettempdir()) / "bolsa-dividendos"),
        help="dónde guardar el Excel bajado (por defecto, una carpeta temporal)",
    )
    parser.add_argument("--hasta", help="bajar el más reciente hasta esta fecha, AAAA-MM-DD (por defecto, hoy)")
    parser.add_argument(
        "--desde",
        help="guardar solo los eventos pagados desde esta fecha, AAAA-MM-DD (por defecto, el 1 de "
        "enero del año anterior al Excel)",
    )
    parser.add_argument("--tabla", default=str(CARPETA_DATOS), help="carpeta de la tabla de la skill")
    args = parser.parse_args(argv)

    try:
        if args.excel:
            excel = Path(args.excel)
        else:
            hasta = date.fromisoformat(args.hasta) if args.hasta else None
            excel = descargar(Path(args.carpeta), hasta)
            print(f"Bajado: {excel}")
        lectura = leer(excel)
    except (RuntimeError, ValueError, OSError) as e:
        parser.error(str(e))

    desde = date.fromisoformat(args.desde) if args.desde else date(lectura.al.year - 1, 1, 1)
    ruta = escribir(lectura, desde, Path(args.tabla))
    print(f"Excel de la Bolsa del {lectura.al:%d-%m-%Y}:")
    for h in lectura.hojas:
        print(
            f"  {h.nombre}: {h.filas} filas = {h.eventos} eventos + {len(h.descartadas)} descartadas "
            f"(sin nemotécnico o fecha de pago); {h.con_monto} con monto por cuota, {h.dudosos} dudosos"
        )
        for texto, n in Counter(h.descartadas).most_common(3):
            print(f"    descartada: {texto}" + (f" (×{n})" if n > 1 else ""))
    clases = Counter(e.clase for e in lectura.eventos)
    print(f"{len(lectura.eventos)} eventos: " + ", ".join(f"{c.lower()} {n}" for c, n in clases.most_common()))

    # La tabla escrita, vuelta a leer, tiene que traer los mismos eventos.
    esperados = Counter(_clave(e) for e in lectura.eventos if e.fecha_pago >= desde)
    leidos = Counter(_clave(e) for e in cargar(ruta).eventos)
    cuadra = esperados == leidos
    print(
        f"{sum(leidos.values())} eventos pagados desde el {desde:%d-%m-%Y} en {ruta}: "
        + ("cuadran con el Excel." if cuadra else f"NO CUADRAN con los {sum(esperados.values())} del Excel.")
    )
    return 0 if cuadra else 1


def _clave(e: ec.Evento) -> tuple:
    return (e.nemotecnico, e.clase, e.moneda, e.monto_por_cuota, e.fecha_limite, e.fecha_pago, e.texto)


if __name__ == "__main__":
    sys.exit(main())
