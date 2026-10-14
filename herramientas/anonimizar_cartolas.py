"""Anonimiza cartolas: reemplaza nombres y RUT de clientes en copias de los PDF y Excel.

Uso (en la carpeta "Interprete Cartolas", donde están "Training Data" y "Example"):

    py -m pip install pymupdf openpyxl
    py anonimizar_cartolas.py

1. La primera vez busca todos los RUT en las cartolas y crea "clientes.csv" con
   cada RUT y el nombre que aparece junto a él. Ábrelo en Excel, corrige los
   nombres y pon "no" en la columna "anonimizar" para los RUT que no son de
   clientes (por ejemplo, el RUT del banco).
   Si después aparecen RUT nuevos (otra cartola, la carpeta Example), se agregan
   al final de "clientes.csv" y el script se detiene para que los revises.
2. Luego busca los demás datos de cada cliente (dirección, comuna, email,
   teléfono y número de cuenta) y crea "otros_datos.csv". Revísalo: pon "no" en
   lo que no sea del cliente, agrega lo que falte y, si un nombre de archivo
   trae una abreviatura del cliente (iniciales, un nombre mal escrito), agrégala
   como tipo "alias". La columna "falso" se llena sola la primera vez.
3. La siguiente vez escribe las copias anonimizadas en la carpeta "anonimizado",
   verifica que no quede ningún dato real y deja un resumen.

Los originales nunca se modifican. "clientes.csv", "otros_datos.csv",
"mapeo_anonimizacion.csv" y "reporte_anonimizacion.txt" tienen datos reales:
quedan solo en este computador, fuera de "anonimizado", y no deben subirse a
ningún lado. Un archivo que no pasa la verificación no se copia a "anonimizado",
así que esa carpeta se puede compartir completa.
"""

import argparse
import csv
import random
import re
import shutil
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import pymupdf
except ImportError:  # versiones antiguas de PyMuPDF
    try:
        import fitz as pymupdf
    except ImportError:
        pymupdf = None

try:
    import openpyxl
    import openpyxl.comments
except ImportError:
    openpyxl = None


CARPETAS_ENTRADA = ("Training Data", "Example")
CARPETA_SALIDA = "anonimizado"
ARCHIVO_CLIENTES = "clientes.csv"
ARCHIVO_MAPEO = "mapeo_anonimizacion.csv"
ARCHIVO_REPORTE = "reporte_anonimizacion.txt"
ARCHIVO_RESUMEN = "RESUMEN.txt"
ARCHIVO_CLAVES = "claves.txt"  # opcional: una clave de PDF por línea
ARCHIVO_OTROS = "otros_datos.csv"

EXT_PDF = {".pdf"}
EXT_EXCEL = {".xlsx", ".xlsm"}
EXT_TEXTO = {".csv", ".txt"}


def aislado(patron, antes, despues=None):
    """`patron` sin un carácter de la clase `antes` justo antes ni de `despues`
    justo después. Equivale a lookarounds negativos, pero se escribe sin ellos
    para que el script pueda viajar como texto por canales que filtran HTML."""
    despues = antes if despues is None else despues
    return r"(?:^|(?<=[^" + antes + r"]))" + patron + r"(?=[^" + despues + r"]|$)"


# RUT con puntos (12.345.678-9) o sin puntos (12345678-9), con o sin espacios
# alrededor del guion. El cuerpo va de 1.000.000 a 99.999.999.
RUT_RE = re.compile(aislado(r"(\d{1,2}(?:\.\d{3}){2}|\d{7,8})(\s*-\s*)([\dkK])", r"\d.", r"\dA-Za-z"))

# Palabras que acompañan a un RUT o a un nombre y que no son parte del nombre.
ETIQUETAS = {
    "RUT", "R.U.T", "R.U.T.", "RUN", "NOMBRE", "NOMBRES", "CLIENTE", "TITULAR",
    "SR", "SR.", "SRA", "SRA.", "SRES", "SRES.", "SENOR", "SENORA", "SENORES",
    "DON", "DONA", "RAZON", "SOCIAL", "N", "NO", "NRO", "NUM",
}

# Palabras cortas que se dejan igual dentro de un nombre de empresa.
CONECTORES = {"DE", "DEL", "LA", "LAS", "EL", "LOS", "Y", "E", "&"}

SUFIJOS_EMPRESA = {
    "SPA", "LTDA", "LTDA.", "LIMITADA", "S.A.", "SA", "S.A", "EIRL", "E.I.R.L.",
    "INVERSIONES", "SOCIEDAD", "COMERCIAL", "ASESORIAS", "INMOBILIARIA",
}

# Dos listas para que los nombres de pila falsos de una persona no mezclen géneros.
# Sin tildes ni eñes: se escriben con la fuente Helvetica del PDF.
NOMBRES_FALSOS = [
    [
        "Ana", "Eva", "Luz", "Ines", "Rosa", "Sara", "Elena", "Paula", "Laura",
        "Marta", "Julia", "Clara", "Sofia", "Lucia", "Andrea", "Camila", "Daniela",
        "Valeria", "Carolina", "Fernanda", "Gabriela", "Francisca", "Constanza",
        "Ada", "Ema", "Alba", "Nora", "Olga", "Irene", "Pilar", "Rocio", "Belen",
        "Silvia", "Teresa", "Lorena", "Paola", "Karina", "Romina", "Cecilia",
        "Natalia", "Viviana", "Pamela", "Monica", "Beatriz", "Antonia", "Josefa",
        "Javiera", "Isidora", "Catalina", "Macarena", "Veronica", "Margarita",
        "Agustina", "Victoria",
    ],
    [
        "Ivo", "Leo", "Luis", "Raul", "Saul", "Hugo", "Ivan", "Pablo", "Pedro",
        "Mario", "Jorge", "Diego", "Tomas", "Andres", "Carlos", "Felipe", "Matias",
        "Rodrigo", "Gonzalo", "Ignacio", "Eduardo", "Cristian", "Sebastian",
        "Francisco", "Maximiliano",
        "Eloy", "Joel", "Noel", "Ruben", "Simon", "Julio", "Oscar", "Ramon",
        "Lucas", "Mateo", "Marcos", "Daniel", "Manuel", "Samuel", "Emilio",
        "Arturo", "Hernan", "Claudio", "Fabian", "Renato", "Ricardo", "Esteban",
        "Joaquin", "Vicente", "Nicolas", "Marcelo", "Mauricio", "Ernesto",
        "Alfredo", "Rolando", "Osvaldo", "Bastian", "Leonardo", "Guillermo",
        "Humberto", "Cristobal", "Benedicto", "Valentin",
    ],
]

APELLIDOS_FALSOS = [
    "Paz", "Lara", "Mena", "Vera", "Rojo", "Soto", "Lagos", "Ortiz", "Pinto",
    "Toro", "Vidal", "Bravo", "Reyes", "Rivas", "Ponce", "Molina", "Robles",
    "Tapia", "Fuentes", "Salinas", "Cabrera", "Aguilera", "Carrasco",
    "Figueroa", "Sepulveda", "Valenzuela", "Riquelme", "Saavedra", "Alvarado",
    "Contreras", "Hernandez", "Villalobos", "Valderrama", "Montenegro",
    "Santibanez", "Echeverria",
    "Mora", "Jara", "Vega", "Diaz", "Lillo", "Moya", "Rojas", "Araya", "Pena",
    "Lobos", "Mella", "Ramos", "Godoy", "Palma", "Ulloa", "Leiva", "Campos",
    "Castro", "Vargas", "Nunez", "Torres", "Flores", "Medina", "Bustos",
    "Ibarra", "Quiroz", "Toledo", "Galvez", "Valdes", "Guzman", "Yanez",
    "Morales", "Herrera", "Cortes", "Navarro", "Caceres", "Vasquez", "Pizarro",
    "Escobar", "Zuniga", "Miranda", "Donoso", "Aravena", "Paredes", "Poblete",
    "Venegas", "Jorquera", "Olivares", "Gallardo", "Orellana", "Sandoval",
    "Castillo", "Espinoza", "Gutierrez", "Cardenas", "Sanhueza", "Cifuentes",
    "Arancibia", "Bustamante", "Rebolledo",
]

# Si el nombre junto a un RUT tiene alguna de estas palabras, probablemente es
# la institución y no un cliente: se propone "no" en clientes.csv.
PALABRAS_INSTITUCION = {
    "BANCO", "BANK", "CORREDORA", "CORREDORES", "AGF", "ADMINISTRADORA", "SEGUROS",
    "BOLSA", "VALORES", "CUSTODIA", "FINANCIERA", "AFP", "FONDO", "FONDOS",
}

PALABRAS_EMPRESA_FALSAS = [
    "Alerce", "Coigue", "Quillay", "Peumo", "Boldo", "Maiten", "Litre",
    "Arrayan", "Canelo", "Lenga", "Rauli", "Ulmo", "Tepa", "Luma",
    "Notro", "Molle", "Radal", "Tineo", "Roble", "Chilco", "Pehuen", "Laurel",
    "Espino", "Queule", "Belloto", "Patagua", "Copihue", "Quintral", "Avellano",
    "Guayacan", "Tamarugo", "Algarrobo",
]

# Palabras de una dirección que se dejan igual (tipo de vía, conectores, unidad).
PALABRAS_DIRECCION = {
    "CALLE", "AVENIDA", "AVDA", "AV", "PASAJE", "PJE", "PSJE", "PAS", "CAMINO",
    "CASA", "DEPTO", "DPTO", "DEPARTAMENTO", "DP", "OFICINA", "OF", "PISO", "P",
    "BLOCK", "TORRE", "PARCELA", "LOTE", "SITIO", "MANZANA", "VILLA", "SAN",
    "SANTA", "A", "B", "C", "D", "CHILE",
} | CONECTORES

CALLES_FALSAS = [
    "Olmos", "Pinos", "Tilos", "Robles", "Sauces", "Cedros", "Lirios", "Aromos",
    "Alamos", "Ulmos", "Lumas", "Tepas", "Notros", "Boldos", "Peumos",
    "Litres", "Fresnos", "Acacias", "Encinas", "Nogales", "Perales", "Mimosas",
    "Dalias", "Retamos", "Camelias", "Naranjos", "Ciruelos", "Duraznos",
    "Jazmines", "Claveles", "Violetas", "Amapolas", "Gardenias", "Magnolias",
    "Abedules", "Laureles", "Copihues", "Arrayanes", "Avellanos", "Araucarias",
    "Hortensias", "Tulipanes", "Limoneros", "Eucaliptos", "Algarrobos",
]

# Comunas que se reconocen al final de una dirección y las que se usan como reemplazo.
COMUNAS_CONOCIDAS = [
    "SAN PEDRO DE LA PAZ", "LO BARNECHEA", "LAS CONDES", "LA REINA", "VITACURA",
    "PROVIDENCIA", "NUNOA", "SANTIAGO", "CONCEPCION", "CHIGUAYANTE", "TALCAHUANO",
    "HUALPEN", "COLINA", "SAN VICENTE", "PENALOLEN", "LA FLORIDA", "MAIPU",
]

COMUNAS_FALSAS = [
    "Hualqui", "Tome", "Penco", "Lota", "Coronel", "Florida", "Cabrero", "Yumbel",
    "Arauco", "Quillon", "Coelemu", "Lebu", "Mulchen", "Nacimiento", "Santa Juana",
    "Macul", "Quilicura", "Pudahuel", "Lampa", "Buin", "Paine", "Talagante",
    "Melipilla", "Pirque", "Til Til", "Curacavi", "Padre Hurtado", "El Monte",
]

DOMINIOS_PUBLICOS = {
    "gmail.com", "hotmail.com", "hotmail.cl", "msn.com", "yahoo.com", "yahoo.es",
    "outlook.com", "live.cl", "live.com", "icloud.com",
}

# Correos de las instituciones: no son datos de clientes y no se tocan.
DOMINIOS_INSTITUCION = {
    "bice.cl", "biceinversiones.cl", "larrainvial.com", "larrainvial.cl", "sura.cl",
    "bny.com", "pershing.com", "inversionessecurity.cl", "security.cl",
}

# Dominios de la asesora (`--dominio-asesora`): en un correo de cliente con ese
# dominio se cambia solo la parte antes de la @. Sin la opción, también el dominio.
DOMINIOS_ASESORA: set[str] = set()


# --------------------------------------------------------------------------
# Utilidades de texto
# --------------------------------------------------------------------------

def sin_acentos(texto):
    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def normalizar(texto):
    """Mayúsculas, sin acentos y con espacios simples, para comparar."""
    return re.sub(r"\s+", " ", sin_acentos(texto).upper()).strip()


def limpiar_token(token):
    return normalizar(token).strip(".,;:()[]\"'")


def digito_verificador(cuerpo):
    suma, factor = 0, 2
    for d in reversed(str(cuerpo)):
        suma += int(d) * factor
        factor = 2 if factor == 7 else factor + 1
    resto = 11 - suma % 11
    return {11: "0", 10: "K"}.get(resto, str(resto))


def cuerpo_rut(texto_cuerpo):
    return int(texto_cuerpo.replace(".", ""))


def formatear_rut(cuerpo, dv, como):
    """Escribe un RUT con el mismo estilo (puntos, separador) que `como`."""
    cuerpo_txt = str(cuerpo)
    if "." in como.group(1):
        cuerpo_txt = f"{cuerpo:,}".replace(",", ".")
    if como.group(3).islower():
        dv = dv.lower()
    return f"{cuerpo_txt}{como.group(2)}{dv}"


def mismo_estilo(original, falso):
    """Copia mayúsculas/minúsculas del original al reemplazo, palabra por palabra."""
    pal_orig, pal_falso = original.split(), falso.split()
    if len(pal_orig) == len(pal_falso) and len(pal_orig) > 1:
        return " ".join(mismo_estilo(o, f) for o, f in zip(pal_orig, pal_falso))
    if limpiar_token(original) == limpiar_token(falso):
        return original  # palabra que no cambia (SpA, Ltda., de, los)
    letras = [c for c in original if c.isalpha()]
    if letras and all(c.isupper() for c in letras):
        return falso.upper()
    if letras and all(c.islower() for c in letras):
        return falso.lower()
    if letras and letras[0].isupper() and all(c.islower() for c in letras[1:]):
        return falso.capitalize()
    return falso


# --------------------------------------------------------------------------
# Clientes y mapeo real -> falso
# --------------------------------------------------------------------------

class Cliente:
    def __init__(self, rut, nombre, tipo):
        self.rut = rut  # int (cuerpo) o None
        self.nombre = nombre.strip()
        self.tipo = tipo
        self.palabras = [t for t in self.nombre.split() if limpiar_token(t)]
        self.tokens = [limpiar_token(t) for t in self.palabras]
        self.falsos = []
        self.rut_falso = None

    @property
    def es_empresa(self):
        return self.tipo == "empresa" or any(t in SUFIJOS_EMPRESA for t in self.tokens)

    def secuencias(self):
        """Órdenes en que el nombre puede aparecer, como listas de índices de token."""
        n = len(self.tokens)
        if n == 0:
            return []
        if self.es_empresa or n == 1:
            return [list(range(n))]
        if n == 2:
            return [[0, 1], [1, 0]]
        nombres = list(range(n - 2))
        apellidos = [n - 2, n - 1]
        primero, paterno = nombres[0], apellidos[0]
        candidatas = [
            nombres + apellidos,
            apellidos + nombres,
            [primero] + apellidos,
            apellidos + [primero],
            nombres + [paterno],
            [paterno] + nombres,
            [primero, paterno],
            [paterno, primero],
            apellidos,
        ]
        if len(nombres) >= 2:
            candidatas.append(nombres)
        unicas = []
        for s in candidatas:
            if s not in unicas:
                unicas.append(s)
        return sorted(unicas, key=len, reverse=True)

    def indices_apellidos(self):
        """Los dos últimos tokens son apellidos (uno solo si el nombre tiene dos)."""
        n = len(self.tokens)
        if self.es_empresa or n < 2:
            return []
        return [n - 1] if n == 2 else [n - 2, n - 1]

    def apellidos_sueltos(self):
        return [i for i in self.indices_apellidos() if len(self.tokens[i]) >= 4]

    def nombres_sueltos(self):
        apellidos = self.indices_apellidos()
        if not apellidos:
            return []
        return [i for i in range(apellidos[0]) if len(self.tokens[i]) >= 3]


def elegir_parecido(lista, largo, usados, rng):
    """Elige una palabra de largo parecido (preferentemente no más larga) que no
    sea real ni se haya usado antes, para que dos clientes no compartan apellido."""
    libres = [p for p in lista if p.upper() not in usados] or list(lista)
    candidatos = sorted(libres, key=lambda p: (len(p) > largo, abs(len(p) - largo)))
    elegido = rng.choice(candidatos[:6])
    usados.add(elegido.upper())
    return elegido


def asignar_falsos(clientes, mapeo_previo, rng):
    ruts_usados = {c.rut for c in clientes if c.rut}
    usados = {t for c in clientes for t in c.tokens}  # palabras reales
    usados |= {normalizar(f) for nombre, _ in mapeo_previo.values() for f in nombre.split()}
    for c in clientes:
        clave = c.rut if c.rut else normalizar(c.nombre)
        if clave in mapeo_previo:
            nombre_falso, rut_falso = mapeo_previo[clave]
            c.falsos = nombre_falso.split()
            c.rut_falso = rut_falso
        if c.rut and not c.rut_falso:
            # mismo largo y mismo primer dígito: un RUT de persona sigue
            # pareciendo de persona y uno de empresa, de empresa
            digitos = len(str(c.rut))
            base = int(str(c.rut)[0]) * 10 ** (digitos - 1)
            while True:
                candidato = rng.randint(base, base + 10 ** (digitos - 1) - 1)
                if candidato not in ruts_usados:
                    break
            ruts_usados.add(candidato)
            c.rut_falso = candidato
        if len(c.falsos) != len(c.tokens):
            c.falsos = []
            # nombre de pila terminado en "a" -> nombres falsos femeninos
            femenino = bool(c.tokens) and c.tokens[0].endswith("A")
            nombres_falsos = NOMBRES_FALSOS[0] if femenino else NOMBRES_FALSOS[1]
            for i, tok in enumerate(c.tokens):
                if len(tok) == 1 and not c.es_empresa:
                    # inicial de un apellido ("Perez Q."): otra letra
                    c.falsos.append(rng.choice([l for l in "ABCDEFGHIJLMNOPRSTUV" if l != tok]))
                    continue
                if c.es_empresa:
                    if tok in SUFIJOS_EMPRESA or tok in CONECTORES:
                        c.falsos.append(c.palabras[i])
                    else:
                        c.falsos.append(elegir_parecido(PALABRAS_EMPRESA_FALSAS, len(tok), usados, rng))
                elif i in c.indices_apellidos():
                    c.falsos.append(elegir_parecido(APELLIDOS_FALSOS, len(tok), usados, rng))
                else:
                    c.falsos.append(elegir_parecido(nombres_falsos, len(tok), usados, rng))


def leer_csv_flexible(ruta):
    texto = ruta.read_text(encoding="utf-8-sig", errors="replace")
    delimitador = ";" if texto.count(";") >= texto.count(",") else ","
    return list(csv.DictReader(texto.splitlines(), delimiter=delimitador))


def escribir_csv(ruta, filas, columnas):
    with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=columnas, delimiter=";")
        w.writeheader()
        for fila in filas:
            w.writerow(fila)


def leer_clientes(ruta):
    clientes = []
    for fila in leer_csv_flexible(ruta):
        fila = {k.strip().lower(): (v or "").strip() for k, v in fila.items() if k}
        if fila.get("anonimizar", "si").lower() in ("no", "n"):
            continue
        rut_txt = fila.get("rut", "")
        m = re.match(r"^\s*([\d.]+)\s*-?\s*[\dkK]?\s*$", rut_txt)
        rut = cuerpo_rut(m.group(1)) if m and m.group(1).replace(".", "") else None
        nombre = fila.get("nombre", "")
        if not rut and not nombre:
            continue
        clientes.append(Cliente(rut, nombre, fila.get("tipo", "").lower()))
    return clientes


def leer_mapeo(ruta):
    mapeo = {}
    if not ruta.exists():
        return mapeo
    for fila in leer_csv_flexible(ruta):
        rut_real = fila.get("rut_real", "").strip()
        clave = int(rut_real) if rut_real.isdigit() else normalizar(fila.get("nombre_real", ""))
        rut_falso = fila.get("rut_falso", "").strip()
        mapeo[clave] = (fila.get("nombre_falso", ""), int(rut_falso) if rut_falso.isdigit() else None)
    return mapeo


def guardar_mapeo(ruta, clientes):
    filas = [{
        "rut_real": c.rut or "",
        "nombre_real": c.nombre,
        "rut_falso": c.rut_falso or "",
        "nombre_falso": " ".join(c.falsos),
    } for c in clientes]
    escribir_csv(ruta, filas, ["rut_real", "nombre_real", "rut_falso", "nombre_falso"])


# --------------------------------------------------------------------------
# Reemplazo en texto plano (Excel, CSV, nombres de archivo)
# --------------------------------------------------------------------------

LETRAS_EQUIVALENTES = {
    "A": "AÁÀÄÂ", "E": "EÉÈËÊ", "I": "IÍÌÏÎ", "O": "OÓÒÖÔ", "U": "UÚÙÜÛ", "N": "NÑ", "C": "CÇ",
}


LETRA = r"\wÁÉÍÓÚÑáéíóúñ"
# Como LETRA pero sin "_": en nombres de archivo el guion bajo separa palabras
# ("1AB123456_MSTAT", "Sura-Agosto-XYZ.pdf").
ALFANUM = r"0-9A-Za-zÁÉÍÓÚÑÜáéíóúñü"


def patron_token(token):
    partes = []
    for c in token:
        grupo = LETRAS_EQUIVALENTES.get(c)
        if grupo:
            partes.append("[" + grupo + grupo.lower() + "]")
        else:
            partes.append(re.escape(c))
    return "".join(partes)


class Reemplazador:
    def __init__(self, clientes):
        self.clientes = clientes
        self.ruts = {c.rut: c for c in clientes if c.rut}
        self.patrones = []  # (regex, cliente, secuencia)
        for c in clientes:
            for seq in c.secuencias():
                piezas = [patron_token(c.tokens[i]) for i in seq]
                sep = r"[\s,]+"
                regex = re.compile(aislado(sep.join(piezas), LETRA), re.IGNORECASE)
                self.patrones.append((regex, c, seq))
            for i in c.apellidos_sueltos():
                regex = re.compile(aislado(patron_token(c.tokens[i]), LETRA), re.IGNORECASE)
                self.patrones.append((regex, c, [i]))
        self.patrones.sort(key=lambda p: len(p[2]), reverse=True)

    def rut_falso_texto(self, m):
        cliente = self.ruts.get(cuerpo_rut(m.group(1)))
        if not cliente:
            return None
        return formatear_rut(cliente.rut_falso, digito_verificador(cliente.rut_falso), m)

    def texto_falso(self, original, cliente, seq):
        palabras = [cliente.falsos[i] for i in seq]
        coma = "," in original
        pal_orig = original.split()
        if coma and len(pal_orig) == len(palabras):
            # la coma queda pegada a la misma palabra: "Asesorias, Inversiones" o "PEREZ, JUAN"
            palabras = [f + "," if o.endswith(",") else f for o, f in zip(pal_orig, palabras)]
            falso = " ".join(palabras)
        elif coma and len(seq) > 1 and cliente.indices_apellidos():
            # "PEREZ GONZALEZ, JUAN": la coma va después de los apellidos.
            n_ap = len([i for i in seq if i in cliente.indices_apellidos()])
            falso = " ".join(palabras[:n_ap]) + ", " + " ".join(palabras[n_ap:])
        else:
            falso = " ".join(palabras)
        return mismo_estilo(original, falso)

    def reemplazar(self, texto):
        cambios = 0

        def sub_rut(m):
            nonlocal cambios
            nuevo = self.rut_falso_texto(m)
            if nuevo is None:
                return m.group(0)
            cambios += 1
            return nuevo

        texto = RUT_RE.sub(sub_rut, texto)
        for regex, cliente, seq in self.patrones:
            def sub_nombre(m, cliente=cliente, seq=seq):
                nonlocal cambios
                cambios += 1
                return self.texto_falso(m.group(0), cliente, seq)
            texto = regex.sub(sub_nombre, texto)
        return texto, cambios

    def reemplazar_sueltos(self, texto):
        """Para nombres de archivo: cambia también una palabra suelta del nombre
        ("SURA-Agosto-Juan"), con el mismo reemplazo que tiene en el nombre completo."""
        if not hasattr(self, "_sueltos"):
            self._sueltos = []
            for c in self.clientes:
                for i, tok in enumerate(c.tokens):
                    if len(tok) >= 3 and tok not in SUFIJOS_EMPRESA and tok not in CONECTORES:
                        regex = re.compile(aislado(patron_token(tok), ALFANUM), re.IGNORECASE)
                        self._sueltos.append((len(tok), regex, c.falsos[i]))
            self._sueltos.sort(key=lambda s: s[0], reverse=True)
        for _, regex, falso in self._sueltos:
            texto = regex.sub(lambda m, falso=falso: mismo_estilo(m.group(0), falso), texto)
        return texto


# --------------------------------------------------------------------------
# Otros datos del cliente: dirección, comuna, email, teléfono, cuenta y alias
# --------------------------------------------------------------------------

# "codigo": identificador corto del cliente dentro del texto (p. ej. sus iniciales como
# código de asesor en Pershing). "alias": igual, pero solo en nombres de archivo.
TIPOS_OTROS = ("direccion", "comuna", "email", "telefono", "cuenta", "codigo", "alias")
COLUMNAS_OTROS = ["tipo", "valor", "cliente", "anonimizar", "falso", "veces", "carpetas"]

EMAIL_RE = re.compile(r"[\w.%+#*-]+@[\w-]+(?:\.[\w-]+)+")
CUENTA_RE = re.compile(r"N[uú]mero de Cuenta:?\s*([A-Z0-9]{2,4}-\d{4,})", re.IGNORECASE)
CODIGO_POSTAL_RE = re.compile(r"^A\d{7}CSF\w*-SD$")  # línea que cierra el bloque de envío (Pershing)
DIRECCION_MAYUS_RE = re.compile(r"^[A-ZÁÉÍÓÚÑ0-9 .]+\s\d+[A-Z0-9 .]*,\s*[A-ZÁÉÍÓÚÑ ]+$")
ETIQUETA_RE = re.compile(r"^\s*[^\d:]{2,30}:\s*$")
SEPARADORES_CUENTA = re.compile(r"[\s\-./]")


def etiqueta_cliente(cliente):
    if cliente is None:
        return ""
    if cliente.rut:
        return f"{cliente.rut:,}".replace(",", ".") + "-" + digito_verificador(cliente.rut)
    return cliente.nombre


def tokens_de(valor):
    return [t for t in (limpiar_token(p) for p in valor.split()) if t]


def iniciales(tokens, indices):
    return "".join(tokens[i][:1] for i in indices if i < len(tokens) and tokens[i])


class OtroDato:
    def __init__(self, fila):
        self.fila = fila
        self.tipo = fila.get("tipo", "").strip().lower()
        self.valor = fila.get("valor", "").strip()
        self.cliente_txt = fila.get("cliente", "").strip()
        self.falso = fila.get("falso", "").strip()
        self.tokens = tokens_de(self.valor)


class OtrosDatos:
    """Valores literales de cada cliente (otros_datos.csv) y su reemplazo falso."""

    def __init__(self, filas, clientes, rng, palabras_reales=()):
        self.rng = rng
        self.por_rut = {c.rut: c for c in clientes if c.rut}
        self.por_nombre = {normalizar(c.nombre): c for c in clientes if c.nombre}
        self.usados = {t for c in clientes for t in c.tokens} | set(palabras_reales)
        self.usados |= {normalizar(f) for c in clientes for f in c.falsos}
        self.numeros, self.palabras, self.comunas, self.dominios = {}, {}, {}, {}
        self.todos = [OtroDato(f) for f in filas]
        self.datos = [d for d in self.todos
                      if d.tipo in TIPOS_OTROS and d.tokens
                      and d.fila.get("anonimizar", "si").strip().lower() not in ("no", "n")]
        self.cuentas = {}
        for d in self.datos:
            for t in d.tokens:
                self.usados.add(t)
            if d.tipo == "cuenta" and d.falso:
                self.cuentas[SEPARADORES_CUENTA.sub("", d.valor)] = SEPARADORES_CUENTA.sub("", d.falso)
        self.generados = 0
        # primero los valores con separadores, para que sus variantes compactas los reutilicen
        pendientes = [d for d in self.datos if not d.falso]
        pendientes.sort(key=lambda d: d.tipo == "cuenta" and not SEPARADORES_CUENTA.search(d.valor))
        for d in pendientes:
            d.falso = self.generar(d)
            d.fila["falso"] = d.falso
            self.generados += 1
        self.datos.sort(key=lambda d: len(d.tokens), reverse=True)
        self.regex_texto = []
        for d in self.datos:
            if d.tipo == "comuna":
                continue  # solo como línea completa en la primera página de un PDF
            patron = r"[\s,]+".join(patron_token(t) for t in d.tokens)
            regex = re.compile(aislado(patron, ALFANUM + "@"), re.IGNORECASE)
            self.regex_texto.append((regex, d))

    # ---- datos falsos ----

    def cliente(self, dato):
        txt = dato.cliente_txt
        m = re.match(r"^\s*([\d.]+)\s*-?\s*[\dkK]?\s*$", txt)
        if m and m.group(1).replace(".", ""):
            return self.por_rut.get(cuerpo_rut(m.group(1)))
        return self.por_nombre.get(normalizar(txt))

    def elegir(self, lista, largo):
        return elegir_parecido(lista, largo, self.usados, self.rng)

    def numero_falso(self, digitos):
        if digitos not in self.numeros:
            while True:
                primero = self.rng.choice("0" if digitos[0] == "0" else "123456789")
                nuevo = primero + "".join(self.rng.choice("0123456789") for _ in digitos[1:])
                if nuevo != digitos or len(digitos) == 1:
                    break
            self.numeros[digitos] = nuevo
        return self.numeros[digitos]

    def palabra_falsa(self, palabra):
        clave = limpiar_token(palabra)
        if clave not in self.palabras:
            self.palabras[clave] = self.elegir(CALLES_FALSAS, len(clave))
        falsa = self.palabras[clave]
        cola = palabra[len(palabra.rstrip(".,;:")):]
        return falsa + cola

    def comuna_falsa(self, comuna):
        clave = normalizar(comuna)
        if clave not in self.comunas:
            self.comunas[clave] = self.elegir(COMUNAS_FALSAS, len(clave))
        return self.comunas[clave]

    def calle_falsa(self, parte):
        palabras = parte.split()
        norm = [limpiar_token(p) for p in palabras]
        comuna = ""
        for c in sorted(COMUNAS_CONOCIDAS, key=len, reverse=True):
            n = len(c.split())
            if len(norm) > n and " ".join(norm[-n:]) == c:
                comuna = self.comuna_falsa(c)
                palabras = palabras[:-n]
                break
        nuevas = []
        for p in palabras:
            t = limpiar_token(p)
            if not t or t in PALABRAS_DIRECCION:
                nuevas.append(p)
            elif re.search(r"\d", p):
                nuevas.append(re.sub(r"\d+", lambda m: self.numero_falso(m.group(0)), p))
            else:
                nuevas.append(self.palabra_falsa(p))
        return " ".join(nuevas + ([comuna] if comuna else []))

    def direccion_falsa(self, valor):
        partes = [p.strip() for p in valor.split(",")]
        nuevas = []
        for k, parte in enumerate(partes):
            if k > 0 and parte and not re.search(r"\d", parte):
                nuevas.append(self.comuna_falsa(parte))  # ", Chiguayante"
            else:
                nuevas.append(self.calle_falsa(parte))
        return ", ".join(nuevas)

    def local_falso(self, local, cliente):
        digitos = sum(ch.isdigit() for ch in local)
        sufijo = "".join(self.rng.choice("0123456789") for _ in range(digitos))
        if cliente and cliente.falsos and not cliente.es_empresa:
            ap = cliente.indices_apellidos()
            if normalizar(local) == iniciales(cliente.tokens, [0] + ap):
                return iniciales([normalizar(f) for f in cliente.falsos], [0] + ap)
            paterno = cliente.falsos[ap[0]] if ap else cliente.falsos[-1]
            return (cliente.falsos[0][:1] + paterno).lower() + sufijo
        if cliente and cliente.falsos:
            propias = [f for f, t in zip(cliente.falsos, cliente.tokens)
                       if t not in SUFIJOS_EMPRESA and t not in CONECTORES]
            return "".join(propias or cliente.falsos).lower() + sufijo
        return self.elegir(CALLES_FALSAS, len(local)).lower() + sufijo

    def dominio_falso(self, dominio, cliente):
        if dominio not in self.dominios:
            base = self.local_falso("", cliente) if cliente else ""
            base = re.sub(r"[^a-z]", "", base.lower()) or self.elegir(CALLES_FALSAS, 6).lower()
            tld = dominio.rsplit(".", 1)[-1]
            nuevo = f"{base}.{tld}"
            while nuevo in self.dominios.values():
                nuevo = f"{base}{self.rng.randint(2, 99)}.{tld}"
            self.dominios[dominio] = nuevo
        return self.dominios[dominio]

    def email_falso(self, valor, cliente):
        local, _, dominio = valor.partition("@")
        prefijo = re.match(r"^[^\w#]*", local).group(0)  # "*" de una nota al pie
        local = local[len(prefijo):]
        if "#" in local or "*" in local:
            # correo que la institución ya enmascaró: x#########@gmail.com
            falso_local = self.local_falso("x", cliente)[:1] + local[1:]
        else:
            falso_local = self.local_falso(local, cliente)
        falso_local = prefijo + falso_local
        dom = dominio.lower()
        if dom in DOMINIOS_PUBLICOS or dom in DOMINIOS_ASESORA:
            falso_dom = dominio
        else:
            falso_dom = self.dominio_falso(dom, cliente)
        return f"{falso_local}@{falso_dom}"

    def telefono_falso(self, valor):
        # se conservan el código de país y de área; cambia el número del abonado
        grupos = list(re.finditer(r"\d+", valor))
        if not grupos:
            return valor
        ultimo = grupos[-1]
        d = ultimo.group(0)
        nuevo = d[0] + "".join(self.rng.choice("0123456789") for _ in d[1:])
        return valor[:ultimo.start()] + nuevo + valor[ultimo.end():]

    def cuenta_falsa(self, valor):
        # cambia cada grupo de 3 o más dígitos; la variante sin guion ("1AB123456")
        # usa el mismo número falso que la con guion ("1AB-123456")
        compacto = SEPARADORES_CUENTA.sub("", valor)
        if compacto in self.cuentas:
            digitos = iter(self.cuentas[compacto])
            return "".join(ch if SEPARADORES_CUENTA.match(ch) else next(digitos) for ch in valor)
        falso = re.sub(r"\d{3,}", lambda m: self.numero_falso(m.group(0)), valor)
        self.cuentas[compacto] = SEPARADORES_CUENTA.sub("", falso)
        return falso

    def alias_falso(self, valor, cliente):
        if not cliente or not cliente.falsos:
            return self.elegir(CALLES_FALSAS, len(valor))
        norm = normalizar(valor)
        ap = cliente.indices_apellidos()
        if norm == iniciales(cliente.tokens, [0] + ap):
            return iniciales([normalizar(f) for f in cliente.falsos], [0] + ap)

        def comun(a, b):
            n = 0
            while n < min(len(a), len(b)) and a[n] == b[n]:
                n += 1
            return n
        i = max(range(len(cliente.tokens)), key=lambda k: comun(norm, cliente.tokens[k]))
        return cliente.falsos[i]

    def generar(self, dato):
        cliente = self.cliente(dato)
        if dato.tipo == "direccion":
            return self.direccion_falsa(dato.valor)
        if dato.tipo == "comuna":
            return self.comuna_falsa(dato.valor)
        if dato.tipo == "email":
            return self.email_falso(dato.valor, cliente)
        if dato.tipo == "telefono":
            return self.telefono_falso(dato.valor)
        if dato.tipo == "cuenta":
            return self.cuenta_falsa(dato.valor)
        return self.alias_falso(dato.valor, cliente)

    # ---- reemplazo ----

    def reemplazar(self, texto, con_alias=False):
        cambios = 0
        for regex, d in self.regex_texto:
            if d.tipo == "alias" and not con_alias:
                continue

            def sub(m, d=d):
                nonlocal cambios
                cambios += 1
                return mismo_estilo(m.group(0), d.falso)
            texto = regex.sub(sub, texto)
        return texto, cambios

    def buscar_en_pagina(self, pagina, es_primera):
        """Como buscar_nombres: secuencias de palabras de una misma línea."""
        hallazgos = []
        for linea in palabras_por_linea(pagina):
            tokens = [limpiar_token(p) for _, p in linea]
            ocupado = [False] * len(linea)
            for d in self.datos:
                if d.tipo == "alias":
                    continue
                n = len(d.tokens)
                if d.tipo == "comuna":
                    # una comuna sola solo se cambia si es la línea completa de la
                    # primera página (bloque de datos del cliente), no en otro texto
                    if not (es_primera and tokens == d.tokens and not any(ocupado)):
                        continue
                for inicio in range(len(tokens) - n + 1):
                    if any(ocupado[inicio:inicio + n]) or tokens[inicio:inicio + n] != d.tokens:
                        continue
                    rect = pymupdf.Rect(linea[inicio][0])
                    for r, _ in linea[inicio + 1:inicio + n]:
                        rect |= r
                    original = " ".join(p for _, p in linea[inicio:inicio + n])
                    primera, final = linea[inicio][1], linea[inicio + n - 1][1]
                    cabeza = primera[:len(primera) - len(primera.lstrip("(\"'¿¡"))]
                    cola = final[len(final.rstrip(".,;:)\"'?!")):]
                    nucleo = original[len(cabeza):len(original) - len(cola)]
                    falso = cabeza + mismo_estilo(nucleo, d.falso)
                    if cola and not falso.endswith(cola):
                        falso += cola
                    hallazgos.append((rect, falso))
                    for k in range(inicio, inicio + n):
                        ocupado[k] = True
        return hallazgos


def leer_otros(ruta):
    filas = []
    for fila in leer_csv_flexible(ruta):
        filas.append({k.strip().lower(): (v or "").strip() for k, v in fila.items() if k})
    return filas


def siguientes_valores(lineas, i, maximo=2):
    """Valores en las líneas que siguen a una etiqueta, hasta la próxima etiqueta."""
    valores = []
    for linea in lineas[i + 1:i + 1 + maximo]:
        if not linea.strip() or ETIQUETA_RE.match(linea):
            break
        valores.append(linea.strip())
    return valores


def direccion_y_comuna(valor):
    """"Calle X 123, San Pedro De La Paz, Concepción, Chile" -> "Calle X 123, San Pedro De La Paz"."""
    partes = [p.strip() for p in valor.split(",") if p.strip()]
    partes = [p for p in partes if normalizar(p) != "CHILE"]
    return ", ".join(partes[:2])


def en_bloque_cliente(lineas, i):
    """True si la etiqueta de la línea i está en el bloque "Cliente:" y no en el del ejecutivo."""
    for linea in reversed(lineas[max(0, i - 8):i]):
        n = normalizar(linea)
        if n.startswith(("EJECUTIVO", "ASESOR", "GESTOR")):
            return False
        if n.startswith("CLIENTE"):
            return True
    return False


def direcciones_bloque_postal(lineas):
    """Bloque de envío de Pershing: nombre, calle, depto y comuna, ciudad, país."""
    fin = next((j for j, l in enumerate(lineas[:20]) if CODIGO_POSTAL_RE.match(l.strip())), None)
    if fin is None:
        return []
    bloque = [l.strip() for l in lineas[:fin] if l.strip()]
    con_numero = [j for j, l in enumerate(bloque) if re.search(r"\d", l)]
    hallados = [("direccion", bloque[j]) for j in con_numero]
    if con_numero:
        ultima = normalizar(bloque[con_numero[-1]])
        tiene_comuna = any(ultima.endswith(" " + c) for c in COMUNAS_CONOCIDAS)
        siguiente = con_numero[-1] + 1
        if not tiene_comuna and siguiente < len(bloque) and normalizar(bloque[siguiente]) in COMUNAS_CONOCIDAS:
            hallados.append(("comuna", bloque[siguiente]))
    return hallados


def cliente_de_archivo(archivo, primera, reemplazador):
    """El cliente de una cartola: RUT en el nombre del archivo o en la primera página,
    o si no hay, el primer nombre de cliente que aparezca en la primera página."""
    for texto in [archivo.name] + primera:
        for m in RUT_RE.finditer(texto):
            c = reemplazador.ruts.get(cuerpo_rut(m.group(1)))
            if c:
                return c
    texto = "\n".join(primera)
    for regex, c, seq in reemplazador.patrones:
        if (len(seq) >= 2 or c.es_empresa) and regex.search(texto):
            return c
    return None


def detectar_otros(archivos, raiz, reemplazador, claves=()):
    hallados = {}

    def agregar(tipo, valor, cliente, carpeta):
        valor = re.sub(r"\s+", " ", valor).strip(" ,;")
        if not tokens_de(valor):
            return
        fila = hallados.setdefault((tipo, normalizar(valor)), {
            "tipo": tipo, "valor": valor, "cliente": etiqueta_cliente(cliente),
            "anonimizar": "si", "falso": "", "veces": 0, "carpetas": set(),
        })
        fila["veces"] += 1
        fila["carpetas"].add(carpeta)

    for archivo in archivos:
        if archivo.suffix.lower() not in EXT_PDF:
            continue
        doc = abrir_pdf(archivo, claves)
        if doc is None:
            continue
        paginas = [p.get_text("text").splitlines() for p in doc]
        doc.close()
        partes = archivo.relative_to(raiz).parts
        carpeta = partes[1] if len(partes) > 2 else partes[0]
        cliente = cliente_de_archivo(archivo, paginas[0] if paginas else [], reemplazador)
        for num, lineas in enumerate(paginas):
            for i, linea in enumerate(lineas):
                for m in EMAIL_RE.finditer(linea):
                    if m.group(0).split("@")[1].lower() not in DOMINIOS_INSTITUCION:
                        agregar("email", m.group(0), cliente, carpeta)
                for m in CUENTA_RE.finditer(linea):
                    agregar("cuenta", m.group(1), cliente, carpeta)
                    agregar("cuenta", m.group(1).replace("-", ""), cliente, carpeta)
                # código del asesor en Pershing: si son las iniciales de un cliente, es él
                m = re.match(r"^\s*N[uú]mero de Identificaci[oó]n:\s*([A-Z]{2,5})\s*$", linea)
                if m:
                    for c in reemplazador.clientes:
                        ap = c.indices_apellidos()
                        if ap and m.group(1) == iniciales(c.tokens, [0] + ap):
                            agregar("codigo", m.group(1), c, carpeta)
                m = re.match(r"^\s*Direcci[oó]n\s*:?\s*(.*)$", linea, re.IGNORECASE)
                if m:
                    valores = [m.group(1)] if m.group(1).strip() else siguientes_valores(lineas, i)
                    for v in valores:
                        agregar("direccion", direccion_y_comuna(v), cliente, carpeta)
                m = re.match(r"^\s*Comuna\s*:?\s*(.*)$", linea, re.IGNORECASE)
                if m and num == 0:
                    valores = [m.group(1)] if m.group(1).strip() else siguientes_valores(lineas, i, 1)
                    for v in valores:
                        agregar("comuna", v, cliente, carpeta)
                m = re.match(r"^\s*Tel[eé]fono\s*:?\s*(.*)$", linea, re.IGNORECASE)
                if m and en_bloque_cliente(lineas, i):
                    v = m.group(1).strip() or (lineas[i + 1].strip() if i + 1 < len(lineas) else "")
                    if sum(ch.isdigit() for ch in v) >= 7:
                        agregar("telefono", v, cliente, carpeta)
            if num == 0:
                for tipo, v in direcciones_bloque_postal(lineas):
                    agregar(tipo, v, cliente, carpeta)
                for linea in lineas[:25]:
                    if DIRECCION_MAYUS_RE.match(linea.strip()):
                        agregar("direccion", linea.strip(), cliente, carpeta)

    filas = sorted(hallados.values(), key=lambda f: (f["cliente"], f["tipo"], f["valor"]))
    for f in filas:
        f["carpetas"] = ", ".join(sorted(f["carpetas"]))
    return filas


# --------------------------------------------------------------------------
# Verificación
# --------------------------------------------------------------------------

class Verificador:
    def __init__(self, clientes, otros=None):
        self.rut_res = []
        self.nombre_res = []
        self.suelto_res = []
        self.otros_res = []  # (tipo, regex)
        self.comunas = set()
        self.alias_res = []
        for d in (otros.datos if otros else []):
            if d.tipo == "comuna":
                self.comunas.add(tuple(d.tokens))
            elif d.tipo == "alias":
                self.alias_res.append(re.compile(aislado(re.escape(" ".join(d.tokens)), "A-Z0-9")))
            else:
                patron = r"[\s,]+".join(re.escape(t) for t in d.tokens)
                self.otros_res.append((d.tipo, re.compile(aislado(patron, r"A-Z0-9@"))))
        # un solo recorrido para todos los RUT (con cientos de clientes, una regex por
        # RUT hace lenta la revisión de archivos grandes)
        self.ruts_reales = {(c.rut, digito_verificador(c.rut)) for c in clientes if c.rut}
        self.rut_res.append(re.compile(aislado(r"(\d{7,8})\s*-?\s*([\dK])", r"\d", r"\dA-Z")))
        for c in clientes:
            for seq in c.secuencias():
                if len(seq) >= 2 or c.es_empresa:
                    self.nombre_res.append(re.compile(
                        aislado(r"[\s,]+".join(re.escape(c.tokens[i]) for i in seq), "A-Z")))
            for i in c.apellidos_sueltos():
                self.nombre_res.append(re.compile(aislado(re.escape(c.tokens[i]), "A-Z")))
            for i in c.nombres_sueltos():
                self.suelto_res.append((c.tokens[i], re.compile(aislado(re.escape(c.tokens[i]), "A-Z"))))

    def revisar(self, texto):
        """Devuelve (problemas, avisos). Problemas impiden copiar el archivo.
        Las comunas se revisan aparte, por posición (anonimizar_pdf)."""
        norm = normalizar(texto)
        sin_puntos = norm.replace(".", "")
        problemas = []
        if any((int(m.group(1)), m.group(2)) in self.ruts_reales
               for r in self.rut_res for m in r.finditer(sin_puntos)):
            problemas.append("queda un RUT real")
        for r in self.nombre_res:
            if r.search(norm):
                problemas.append("queda un nombre real")
                break
        tipos = sorted({tipo for tipo, r in self.otros_res if r.search(norm)})
        if tipos:
            problemas.append("queda un dato real: " + ", ".join(tipos))
        avisos = sorted({tok for tok, r in self.suelto_res if r.search(norm)})
        return problemas, avisos

    def revisar_nombre_archivo(self, ruta):
        """Para nombres de archivo también cuentan un nombre de pila o un alias sueltos."""
        texto = re.sub(r"[_\-]+", " ", str(ruta))
        problemas, avisos = self.revisar(texto)
        if avisos:
            problemas.append("el nombre del archivo tiene un nombre de pila real")
        norm = normalizar(texto)
        if any(r.search(norm) for r in self.alias_res):
            problemas.append("el nombre del archivo tiene un alias real")
        return problemas


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def texto_legible(texto):
    utiles = [c for c in texto if not c.isspace()]
    if not utiles:
        return True
    raros = sum(1 for c in utiles if unicodedata.category(c).startswith("C") or c == "\ufffd")
    return raros / len(utiles) < 0.05


def pagina_escaneada(pagina):
    texto = pagina.get_text("text").strip()
    if len(texto) >= 50:
        return False
    area_pagina = abs(pagina.rect)
    area_imagenes = 0
    for info in pagina.get_image_info():
        area_imagenes += abs(pymupdf.Rect(info["bbox"]) & pagina.rect)
    return area_imagenes > 0.5 * area_pagina


def estilo_en(bloques, rect):
    """Tamaño, color y negrita del texto que cae en `rect`."""
    mejor = None
    for bloque in bloques:
        for linea in bloque.get("lines", []):
            for span in linea["spans"]:
                inter = pymupdf.Rect(span["bbox"]) & rect
                if inter.is_empty:
                    continue
                area = abs(inter)
                if mejor is None or area > mejor[0]:
                    mejor = (area, span, linea["dir"])
    if mejor is None:
        return 9, (0, 0, 0), False, (rect.x0, rect.y1), (1, 0)
    _, span, direccion = mejor
    color = span["color"]
    rgb = ((color >> 16) & 255) / 255, ((color >> 8) & 255) / 255, (color & 255) / 255
    negrita = bool(span["flags"] & 16) or "bold" in span["font"].lower()
    return span["size"], rgb, negrita, span["origin"], (round(direccion[0]), round(direccion[1]))


# Dirección del texto -> (rotación para insert_text, cómo ubicar el inicio)
ROTACIONES = {
    (1, 0): 0,     # horizontal normal
    (0, -1): 90,   # hacia arriba
    (-1, 0): 180,  # cabeza abajo
    (0, 1): 270,   # hacia abajo
}


def punto_inicio(rect, origen, direccion):
    """Dónde empieza a escribirse el reemplazo, sobre la línea base original."""
    return {
        (1, 0): pymupdf.Point(rect.x0, origen[1]),
        (0, -1): pymupdf.Point(origen[0], rect.y1),
        (-1, 0): pymupdf.Point(rect.x1, origen[1]),
        (0, 1): pymupdf.Point(origen[0], rect.y0),
    }[direccion]


def franja_central(rect, direccion):
    """La redacción borra todo carácter que toque el rectángulo; se usa solo la
    franja central (a lo largo del texto) para no borrar letras de líneas vecinas."""
    if direccion[1] == 0:
        alto = rect.height
        return pymupdf.Rect(rect.x0 + 0.3, rect.y0 + 0.3 * alto, rect.x1 - 0.3, rect.y1 - 0.3 * alto)
    ancho = rect.width
    return pymupdf.Rect(rect.x0 + 0.3 * ancho, rect.y0 + 0.3, rect.x1 - 0.3 * ancho, rect.y1 - 0.3)


def franja_segura(rect, direccion, otras):
    """Como franja_central, pero si otra palabra se monta sobre el dato (p. ej. un
    texto blanco escondido detrás de otro párrafo), usa la parte del alto del dato
    que ninguna otra palabra toca. Devuelve None si no queda ninguna parte libre."""
    horizontal = direccion[1] == 0
    lo, hi = (rect.y0, rect.y1) if horizontal else (rect.x0, rect.x1)
    a0, a1 = (rect.x0 + 0.3, rect.x1 - 0.3) if horizontal else (rect.y0 + 0.3, rect.y1 - 0.3)
    ocupados = []
    for o in otras:
        o0, o1 = (o.x0, o.x1) if horizontal else (o.y0, o.y1)
        t0, t1 = (o.y0, o.y1) if horizontal else (o.x0, o.x1)
        if o1 > a0 and o0 < a1 and t1 > lo and t0 < hi:
            ocupados.append((max(t0, lo), min(t1, hi)))
    libres, pos = [], lo
    for t0, t1 in sorted(ocupados):
        if t0 > pos:
            libres.append((pos, t0))
        pos = max(pos, t1)
    if pos < hi:
        libres.append((pos, hi))

    def armar(b0, b1):
        return pymupdf.Rect(a0, b0, a1, b1) if horizontal else pymupdf.Rect(b0, a0, b1, a1)
    c0, c1 = lo + 0.3 * (hi - lo), hi - 0.3 * (hi - lo)
    if any(l0 <= c0 and c1 <= l1 for l0, l1 in libres):
        return armar(c0, c1)
    mejor = max(libres, key=lambda l: l[1] - l[0], default=None)
    if mejor is None or mejor[1] - mejor[0] < 0.6:
        return None
    margen = min(0.2, (mejor[1] - mejor[0]) / 4)
    return armar(mejor[0] + margen, mejor[1] - margen)


def opciones_redaccion():
    opciones = {"images": pymupdf.PDF_REDACT_IMAGE_NONE}
    if hasattr(pymupdf, "PDF_REDACT_LINE_ART_NONE"):
        opciones["graphics"] = pymupdf.PDF_REDACT_LINE_ART_NONE
    return opciones


def borra_solo_el_dato(doc, numero, franja, rect):
    """Prueba la redacción en una copia de la página: True si borra todas las
    palabras del dato y ninguna otra (ni deja pedazos de palabras)."""
    copia = pymupdf.open()
    copia.insert_pdf(doc, from_page=numero, to_page=numero)
    pagina = copia[0]
    antes = contar_palabras(pagina)
    objetivo = palabras_en(pagina, [rect])
    pagina.add_redact_annot(franja, fill=None)
    pagina.apply_redactions(**opciones_redaccion())
    despues = contar_palabras(pagina)
    copia.close()
    return bool(objetivo) and antes - despues == objetivo and not (despues - antes)


def elegir_franja(doc, numero, rect, direccion, otras):
    """La franja central sirve si ninguna otra palabra la toca. Si alguna se monta
    sobre el dato, las cajas de get_text no bastan para saber qué borrará la
    redacción (son más altas que las que usa MuPDF): se prueban franjas delgadas a
    distintas alturas en una copia de la página y se usa la primera que borra solo
    el dato. None si ninguna sirve."""
    central = franja_central(rect, direccion)
    if not any(abs(central & o) > 0 for o in otras):
        return central
    candidatas = [franja_segura(rect, direccion, otras)]
    horizontal = direccion[1] == 0
    lo, hi = (rect.y0, rect.y1) if horizontal else (rect.x0, rect.x1)
    a0, a1 = (rect.x0 + 0.3, rect.x1 - 0.3) if horizontal else (rect.y0 + 0.3, rect.y1 - 0.3)
    for fraccion in (0.5, 0.4, 0.6, 0.3, 0.7, 0.2, 0.8, 0.12, 0.88):
        b = lo + fraccion * (hi - lo)
        candidatas.append(pymupdf.Rect(a0, b - 0.1, a1, b + 0.1) if horizontal
                          else pymupdf.Rect(b - 0.1, a0, b + 0.1, a1))
    for franja in candidatas:
        if franja is not None and borra_solo_el_dato(doc, numero, franja, rect):
            return franja
    return None


def dentro_de(rect_palabra, rects):
    """La palabra está (casi) entera dentro de uno de los rectángulos reemplazados."""
    area = abs(rect_palabra)
    return area > 0 and any(abs(rect_palabra & r) >= 0.8 * area for r in rects)


def lineas_comuna(pagina, comunas):
    """Rectángulos de las líneas de la página que son solo una comuna de cliente."""
    rects = []
    for linea in palabras_por_linea(pagina):
        if tuple(t for t in (limpiar_token(p) for _, p in linea) if t) in comunas:
            r = pymupdf.Rect(linea[0][0])
            for otro, _ in linea[1:]:
                r |= otro
            rects.append(r)
    return rects


def palabras_por_linea(pagina):
    lineas = {}
    for x0, y0, x1, y1, palabra, b, l, w in pagina.get_text("words"):
        lineas.setdefault((b, l), []).append((pymupdf.Rect(x0, y0, x1, y1), palabra))
    return list(lineas.values())


def buscar_nombres(pagina, reemplazador):
    """Busca secuencias de nombre en palabras consecutivas de una misma línea."""
    hallazgos = []
    for linea in palabras_por_linea(pagina):
        tokens = [limpiar_token(p) for _, p in linea]
        ocupado = [False] * len(linea)
        for _, cliente, seq in reemplazador.patrones:
            buscado = [cliente.tokens[i] for i in seq]
            n = len(buscado)
            for inicio in range(len(tokens) - n + 1):
                if any(ocupado[inicio:inicio + n]):
                    continue
                if tokens[inicio:inicio + n] == buscado:
                    rect = pymupdf.Rect(linea[inicio][0])
                    for r, _ in linea[inicio + 1:inicio + n]:
                        rect |= r
                    original = " ".join(p for _, p in linea[inicio:inicio + n])
                    falso = reemplazador.texto_falso(original, cliente, seq)
                    # conserva la puntuación pegada a las palabras, p. ej. "(PEREZ," o "PEREZ."
                    primera, final = linea[inicio][1], linea[inicio + n - 1][1]
                    cabeza = primera[:len(primera) - len(primera.lstrip("(\"'¿¡"))]
                    cola = final[len(final.rstrip(".,;:)\"'?!")):]
                    falso = cabeza + falso
                    if cola and not falso.endswith(cola):
                        falso += cola
                    inicial = inicial_materna(cliente, seq, linea, inicio + n, ocupado)
                    if inicial:
                        rect |= linea[inicio + n][0]
                        falso += " " + inicial
                        ocupado[inicio + n] = True
                    hallazgos.append((rect, falso))
                    for k in range(inicio, inicio + n):
                        ocupado[k] = True
    return hallazgos


def inicial_materna(cliente, seq, linea, pos, ocupado):
    """En "Perez S." la inicial del apellido materno también identifica: si la
    secuencia termina en el apellido paterno y la palabra siguiente es la inicial
    del materno, devuelve la inicial falsa ("P."). Si no, None."""
    apellidos = cliente.indices_apellidos()
    if len(apellidos) != 2 or seq[-1] != apellidos[0] or apellidos[1] in seq:
        return None
    if pos >= len(linea) or ocupado[pos]:
        return None
    palabra = linea[pos][1]
    if not re.fullmatch(r"[^\W\d_]\.?", palabra):
        return None
    if limpiar_token(palabra) != cliente.tokens[apellidos[1]][:1]:
        return None
    return limpiar_token(cliente.falsos[apellidos[1]])[:1] + ("." if palabra.endswith(".") else "")


def buscar_ruts(pagina, reemplazador):
    hallazgos = []
    texto = pagina.get_text("text")
    vistos = set()
    for m in sorted(RUT_RE.finditer(texto), key=lambda m: len(m.group(0)), reverse=True):
        nuevo = reemplazador.rut_falso_texto(m)
        if nuevo is None or m.group(0) in vistos:
            continue
        vistos.add(m.group(0))
        for rect in pagina.search_for(m.group(0)):
            hallazgos.append((rect, nuevo))
    return hallazgos


def claves_de_rut(cuerpo):
    """Los bancos suelen proteger la cartola con el RUT del cliente o parte de él."""
    r, dv = str(cuerpo), digito_verificador(cuerpo)
    return [r, r + dv, r + "-" + dv, f"{cuerpo:,}".replace(",", ".") + "-" + dv, r[:4], r[:5], r[:6]]


def leer_claves_extra(raiz):
    ruta = raiz / ARCHIVO_CLAVES
    if not ruta.exists():
        return []
    return [l.strip() for l in ruta.read_text(encoding="utf-8-sig", errors="replace").splitlines() if l.strip()]


def claves_posibles(ruts, extra=()):
    claves = list(extra)
    for cuerpo in ruts:
        claves += claves_de_rut(cuerpo)
    return list(dict.fromkeys(claves))


def abrir_pdf(ruta, claves=()):
    """Abre el PDF, probando las claves si tiene. Devuelve None si ninguna sirve.

    Ojo: no se consulta needs_pass ni is_encrypted después de authenticate();
    en algunas versiones de PyMuPDF eso deja el contenido sin descifrar.
    """
    doc = pymupdf.open(ruta)
    if not doc.needs_pass:
        return doc
    for clave in claves:
        if doc.authenticate(clave):
            return doc
    doc.close()
    return None


def largo_texto(texto):
    return len(re.sub(r"\s+", "", texto))


def contar_palabras(pagina):
    return Counter(t for t in (limpiar_token(w[4]) for w in pagina.get_text("words")) if t)


def palabras_en(pagina, rects):
    """Palabras que quedan dentro de alguno de los rectángulos reemplazados."""
    dentro = Counter()
    for x0, y0, x1, y1, palabra, *_ in pagina.get_text("words"):
        if dentro_de(pymupdf.Rect(x0, y0, x1, y1), rects):
            t = limpiar_token(palabra)
            if t:
                dentro[t] += 1
    return dentro


def tamano_que_cabe(texto, tamano, fuente, rect, direccion):
    """Achica la letra (hasta 70 %) si el reemplazo no cabe donde estaba el original,
    para que no se monte sobre la palabra siguiente."""
    ancho = rect.width if direccion[1] == 0 else rect.height
    largo = pymupdf.get_text_length(texto, fontname=fuente, fontsize=tamano)
    if largo <= ancho or largo == 0:
        return tamano
    return max(tamano * ancho / largo, tamano * 0.7)


def valores_internos(doc):
    """Los diccionarios de todos los objetos del PDF (metadatos, enlaces, formularios,
    marcadores...) y los streams de metadatos y adjuntos, para buscar datos reales
    fuera del texto de las páginas. Las fuentes e imágenes no se revisan: son
    binarias y pueden contener cualquier secuencia de letras."""
    partes = []
    for xref in range(1, doc.xref_length()):
        try:
            objeto = doc.xref_object(xref, compressed=True)
            partes.append(objeto)
            if doc.xref_is_stream(xref) and re.search(r"/Type\s*/(Metadata|EmbeddedFile)", objeto):
                partes.append((doc.xref_stream(xref) or b"").decode("utf-8", "replace"))
        except Exception:
            continue
    return "\n".join(partes)


def texto_sin_recorte(pagina):
    """Texto de la página incluido el que queda fuera del área visible."""
    flags = pymupdf.TEXTFLAGS_TEXT & ~getattr(pymupdf, "TEXT_MEDIABOX_CLIP", 0)
    return pagina.get_text("text", flags=flags)


def anonimizar_pdf(origen, destino, reemplazador, verificador, claves=(), otros=None):
    doc = abrir_pdf(origen, claves)
    if doc is None:
        return 0, [f"el PDF tiene clave y no es el RUT de ningún cliente ni está en {ARCHIVO_CLAVES}"], []
    texto_original = largo_texto("".join(p.get_text("text") for p in doc))
    if texto_original == 0:
        doc.close()
        return 0, ["el PDF no tiene texto que se pueda leer (¿escaneado?)"], []
    problemas, avisos, cambios = [], [], 0
    # líneas de la primera página que son solo la comuna del cliente ("Comuna:" de BICE)
    comunas = getattr(verificador, "comunas", set())
    lugares_comuna = lineas_comuna(doc[0], comunas) if doc.page_count and comunas else []
    for pagina in doc:
        for campo in pagina.widgets() or []:
            if isinstance(campo.field_value, str):
                nuevo, n = reemplazador.reemplazar(campo.field_value)
                if otros:
                    nuevo, m = otros.reemplazar(nuevo)
                    n += m
                if n:
                    campo.field_value = nuevo
                    campo.update()
                    cambios += n
        for anotacion in pagina.annots() or []:
            info = anotacion.info
            nuevo, n = reemplazador.reemplazar(info.get("content", ""))
            titulo, m = reemplazador.reemplazar(info.get("title", ""))
            if otros:
                nuevo, n2 = otros.reemplazar(nuevo)
                titulo, m2 = otros.reemplazar(titulo)
                n, m = n + n2, m + m2
            if n or m:
                anotacion.set_info(content=nuevo, title=titulo)
                anotacion.update()
                cambios += n + m
        if pagina_escaneada(pagina):
            problemas.append(f"la página {pagina.number + 1} parece escaneada (imagen sin texto)")
            continue
        if not texto_legible(pagina.get_text("text")):
            problemas.append(f"el texto de la página {pagina.number + 1} no se puede leer bien")
            continue
        hallazgos = []
        candidatos = buscar_ruts(pagina, reemplazador)
        if otros:
            candidatos += otros.buscar_en_pagina(pagina, pagina.number == 0)
        candidatos += buscar_nombres(pagina, reemplazador)
        for rect, nuevo in candidatos:
            if any(abs(rect & r) > 0.3 * abs(rect) for r, _ in hallazgos):
                continue
            hallazgos.append((rect, nuevo))
        if not hallazgos:
            continue
        bloques = pagina.get_text("dict")["blocks"]
        estilos = [estilo_en(bloques, rect) for rect, _ in hallazgos]
        if any(e[4] not in ROTACIONES for e in estilos):
            problemas.append(f"la página {pagina.number + 1} tiene un nombre o RUT en texto inclinado")
            continue
        antes = contar_palabras(pagina)
        rects = [r for r, _ in hallazgos]
        reemplazadas = palabras_en(pagina, rects)
        otras = [pymupdf.Rect(w[:4]) for w in pagina.get_text("words")
                 if not dentro_de(pymupdf.Rect(w[:4]), rects)]
        franjas = [elegir_franja(doc, pagina.number, rect, estilo[4], otras)
                   for rect, estilo in zip(rects, estilos)]
        if any(f is None for f in franjas):
            problemas.append(f"la página {pagina.number + 1} tiene un dato montado sobre otra línea "
                             f"y no se puede borrar sin tocarla")
            continue
        for franja in franjas:
            pagina.add_redact_annot(franja, fill=None)
        pagina.apply_redactions(**opciones_redaccion())
        for (rect, nuevo), (tamano, color, negrita, origen, direccion) in zip(hallazgos, estilos):
            fuente = "hebo" if negrita else "helv"
            pagina.insert_text(
                punto_inicio(rect, origen, direccion), nuevo,
                fontsize=tamano_que_cabe(nuevo, tamano, fuente, rect, direccion),
                fontname=fuente, color=color,
                rotate=ROTACIONES[direccion],
                # texto blanco = escondido detrás de otro: va debajo, como el original,
                # para no borronear el texto visible que tiene encima
                overlay=min(color) < 0.95,
            )
            cambios += 1
        # Ninguna palabra fuera de los datos reemplazados puede desaparecer: si la
        # redacción borró letras de una línea vecina, el montaje no es confiable.
        perdidas = (antes - contar_palabras(pagina)) - reemplazadas
        if perdidas:
            problemas.append(f"la página {pagina.number + 1} perdió {sum(perdidas.values())} "
                             f"palabras que no eran datos del cliente")
    if problemas:
        doc.close()
        return cambios, problemas, avisos

    doc.set_metadata({})
    try:
        doc.del_xml_metadata()
    except Exception:
        pass
    toc = doc.get_toc(simple=True)
    if toc:
        doc.set_toc([[nivel, reemplazador.reemplazar(titulo)[0], pag] for nivel, titulo, pag in toc])

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino, garbage=4, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE)
    doc.close()

    verificado = pymupdf.open(destino)
    partes = []
    for p in verificado:
        partes.append(texto_sin_recorte(p))
        partes += [str(w.field_value) for w in (p.widgets() or []) if w.field_value]
        partes += [a.info.get("content", "") + " " + a.info.get("title", "") for a in (p.annots() or [])]
        partes += [enlace.get("uri", "") or "" for enlace in p.get_links()]
    texto = "\n".join(partes)
    texto += "\n" + " ".join(str(v) for v in (verificado.metadata or {}).values() if v)
    quedan_comunas = [r for r in (lineas_comuna(verificado[0], comunas) if lugares_comuna else [])
                      if any(abs(r & l) > 0.5 * abs(r) for l in lugares_comuna)]
    texto_final = largo_texto("".join(p.get_text("text") for p in verificado))
    internos = valores_internos(verificado)
    verificado.close()
    problemas, avisos = verificador.revisar(texto)
    if quedan_comunas:
        problemas.append("queda la comuna real del cliente")
    problemas_internos, _ = verificador.revisar(internos)
    if problemas_internos:
        problemas.append("dentro del archivo (fuera del texto visible) " + "; ".join(problemas_internos))
    # Los reemplazos tienen largo parecido; si se pierde mucho texto, algo salió mal
    # y la verificación no sería confiable.
    if texto_final < 0.9 * texto_original:
        problemas.append(f"la copia perdió texto ({texto_final} de {texto_original} caracteres)")
    if problemas:
        destino.unlink()
    return cambios, problemas, avisos


# --------------------------------------------------------------------------
# Excel y texto
# --------------------------------------------------------------------------

def reemplazar_todo(texto, reemplazador, otros=None):
    """RUT y nombres, y después los otros datos del cliente."""
    texto, n = reemplazador.reemplazar(texto)
    if otros:
        texto, m = otros.reemplazar(texto)
        n += m
    return texto, n


def anonimizar_excel(origen, destino, reemplazador, verificador, otros=None):
    libro = openpyxl.load_workbook(origen, keep_vba=origen.suffix.lower() == ".xlsm")
    cambios = 0
    for hoja in libro.worksheets:
        titulo, n = reemplazar_todo(hoja.title, reemplazador, otros)
        if n:
            hoja.title = titulo[:31]
            cambios += n
        columnas_rut, columnas_dv = set(), set()
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.comment is not None:
                    texto, n = reemplazar_todo(celda.comment.text, reemplazador, otros)
                    autor, m = reemplazar_todo(celda.comment.author or "", reemplazador, otros)
                    if n or m:
                        celda.comment = openpyxl.comments.Comment(texto, autor)
                        cambios += n + m
                v = celda.value
                if isinstance(v, str):
                    if re.search(r"\bR\.?U\.?[TN]\.?\b", sin_acentos(v).upper()):
                        columnas_rut.add(celda.column)
                    if re.fullmatch(r"\s*D\.?\s*V\.?\s*", v.upper()):
                        columnas_dv.add(celda.column)
                    nuevo, n = reemplazar_todo(v, reemplazador, otros)
                    if n:
                        celda.value = nuevo
                        cambios += n
                elif isinstance(v, int) and celda.column in columnas_rut and v in reemplazador.ruts:
                    # RUT guardado como número en una columna "RUT"
                    falso = reemplazador.ruts[v].rut_falso
                    celda.value = falso
                    cambios += 1
                    if celda.column + 1 in columnas_dv:
                        dv = hoja.cell(row=celda.row, column=celda.column + 1)
                        dv_falso = digito_verificador(falso)
                        dv.value = int(dv_falso) if isinstance(dv.value, int) and dv_falso.isdigit() else dv_falso
    libro.properties.creator = None
    libro.properties.lastModifiedBy = None
    libro.properties.title = None
    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)

    verificado = openpyxl.load_workbook(destino)
    textos, ruts_numericos = [], False
    for hoja in verificado.worksheets:
        textos.append(hoja.title)
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.value is not None:
                    textos.append(str(celda.value))
                    ruts_numericos |= isinstance(celda.value, int) and celda.value in reemplazador.ruts
                if celda.comment is not None:
                    textos.append(f"{celda.comment.text} {celda.comment.author or ''}")
    verificado.close()
    problemas, avisos = verificador.revisar("\n".join(textos))
    if ruts_numericos:
        problemas.append("queda un RUT real guardado como número")
    if problemas:
        destino.unlink()
    return cambios, problemas, avisos


def anonimizar_texto(origen, destino, reemplazador, verificador, otros=None):
    contenido = origen.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    nuevo, cambios = reemplazar_todo(texto, reemplazador, otros)
    problemas, avisos = verificador.revisar(nuevo)
    if codificacion == "utf-8-sig" and not contenido.startswith(b"\xef\xbb\xbf"):
        codificacion = "utf-8"  # sin esto la copia gana un BOM que el original no tenía
    if not problemas:
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(nuevo.encode(codificacion))
    return cambios, problemas, avisos


# --------------------------------------------------------------------------
# Detección inicial de clientes
# --------------------------------------------------------------------------

def parece_nombre(texto):
    tokens = [t for t in texto.replace(",", " ").split() if t]
    limpios = [limpiar_token(t) for t in tokens]
    limpios = [t for t in limpios if t and t not in ETIQUETAS]
    if not 2 <= len(limpios) <= 7:
        return None
    if not all(re.fullmatch(r"[A-Z.&'-]+", t) for t in limpios):
        return None
    return " ".join(t for t in tokens if limpiar_token(t) not in ETIQUETAS and limpiar_token(t))


class PdfConClave(Exception):
    pass


def textos_de(archivo, claves=()):
    """Líneas de texto de un archivo, para buscar RUT y nombres."""
    ext = archivo.suffix.lower()
    if ext in EXT_PDF:
        doc = abrir_pdf(archivo, claves)
        if doc is None:
            raise PdfConClave()
        lineas = [l for p in doc for l in p.get_text("text").splitlines()]
        doc.close()
        return lineas
    if ext in EXT_EXCEL:
        libro = openpyxl.load_workbook(archivo, read_only=True)
        lineas = []
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows(values_only=True):
                lineas.append("  ".join(str(v) for v in fila if v is not None))
        libro.close()
        return lineas
    if ext in EXT_TEXTO:
        return archivo.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    return []


def detectar_clientes(archivos, raiz, claves_extra=()):
    encontrados = {}
    con_clave = []

    def revisar(archivo, claves):
        try:
            lineas = textos_de(archivo, claves)
        except PdfConClave:
            return False
        except Exception as e:
            print(f"  No pude leer {archivo.relative_to(raiz)}: {e}")
            return True
        partes = archivo.relative_to(raiz).parts
        carpeta = partes[1] if len(partes) > 2 else partes[0]
        for i, linea in enumerate(lineas):
            for m in RUT_RE.finditer(linea):
                cuerpo = cuerpo_rut(m.group(1))
                dv = m.group(3).upper()
                if digito_verificador(cuerpo) != dv:
                    continue  # no es un RUT válido (probablemente otro número)
                info = encontrados.setdefault(cuerpo, {"dv": dv, "nombres": {}, "veces": 0, "carpetas": set()})
                info["veces"] += 1
                info["carpetas"].add(carpeta)
                resto = (linea[:m.start()] + " " + linea[m.end():]).strip(" :-")
                candidatos = [resto] + [lineas[j] for j in (i - 1, i + 1) if 0 <= j < len(lineas)]
                for cand in candidatos:
                    nombre = parece_nombre(cand)
                    if nombre:
                        info["nombres"][nombre] = info["nombres"].get(nombre, 0) + 1
                        break
        return True

    for archivo in archivos:
        if not revisar(archivo, claves_extra):
            con_clave.append(archivo)
    # Los PDF con clave se reintentan con los RUT encontrados en los demás archivos.
    sin_abrir = [a for a in con_clave if not revisar(a, claves_posibles(encontrados, claves_extra))]
    if sin_abrir:
        print(f"  {len(sin_abrir)} PDF tienen una clave que no es ninguno de los RUT encontrados.")
        print(f"  Si sabes la clave, escríbela en {ARCHIVO_CLAVES} (una por línea) o agrega el RUT del cliente a {ARCHIVO_CLIENTES}.")

    filas = []
    for cuerpo, info in sorted(encontrados.items(), key=lambda kv: -kv[1]["veces"]):
        nombre = max(info["nombres"], key=info["nombres"].get) if info["nombres"] else ""
        institucion = any(limpiar_token(t) in PALABRAS_INSTITUCION for t in nombre.split())
        filas.append({
            "rut": f"{cuerpo:,}".replace(",", ".") + "-" + info["dv"],
            "nombre": nombre,
            "tipo": "persona" if cuerpo < 50_000_000 else "empresa",
            "anonimizar": "no" if institucion else "si",
            "veces": info["veces"],
            "carpetas": ", ".join(sorted(info["carpetas"])),
        })
    return filas


# --------------------------------------------------------------------------
# Programa principal
# --------------------------------------------------------------------------

def agregar_ruts_nuevos(ruta_clientes, archivos, raiz, claves_extra=()):
    """Agrega al final de clientes.csv los RUT que aparecen en los archivos y que
    todavía no están en la lista (ni siquiera con "no"). Devuelve las filas nuevas."""
    filas = leer_csv_flexible(ruta_clientes)
    columnas = ["rut", "nombre", "tipo", "anonimizar", "veces", "carpetas"]
    for fila in filas:
        columnas += [k for k in fila if k and k not in columnas]
    conocidos = set()
    for fila in filas:
        m = re.match(r"^\s*([\d.]+)", fila.get("rut", "") or "")
        if m and m.group(1).replace(".", ""):
            conocidos.add(cuerpo_rut(m.group(1)))
    nuevos = [f for f in detectar_clientes(archivos, raiz, claves_extra)
              if cuerpo_rut(f["rut"].split("-")[0]) not in conocidos]
    if nuevos:
        escribir_csv(ruta_clientes, [{k: f.get(k, "") for k in columnas} for f in filas + nuevos], columnas)
    return nuevos


def nombre_falso(parte, reemplazador, otros):
    """Nombre de archivo o carpeta: RUT, nombres, otros datos, alias y palabras sueltas."""
    nuevo = reemplazador.reemplazar(parte)[0]
    if otros:
        nuevo = otros.reemplazar(nuevo, con_alias=True)[0]
    return reemplazador.reemplazar_sueltos(nuevo)


def listar_archivos(raiz):
    archivos = []
    for nombre in CARPETAS_ENTRADA:
        # "Example" también acepta "Example Export" y similares
        carpeta = next((p for p in raiz.iterdir() if p.is_dir() and (
            p.name.lower() == nombre.lower() or p.name.lower().startswith(nombre.lower() + " "))), None)
        if carpeta is None:
            print(f"Aviso: no encontré la carpeta \"{nombre}\" en {raiz}")
            continue
        archivos.extend(sorted(p for p in carpeta.rglob("*") if p.is_file()))
    return archivos


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--carpeta", default=None,
                        help="carpeta que contiene Training Data y Example (por defecto, la de este script "
                             "o la carpeta actual)")
    parser.add_argument("--semilla", type=int, default=None, help="semilla para los datos falsos (pruebas)")
    parser.add_argument("--dominio-asesora", action="append", default=[], metavar="DOMINIO",
                        help="dominio de correo de la asesora (se puede repetir): en los correos de clientes "
                             "con ese dominio se cambia solo la parte antes de la @")
    args = parser.parse_args()
    DOMINIOS_ASESORA.update(d.strip().lstrip("@").lower() for d in args.dominio_asesora if d.strip())

    if pymupdf is None or openpyxl is None:
        print("Faltan librerías. Instálalas con:\n    py -m pip install pymupdf openpyxl")
        return 1

    if args.carpeta:
        raiz = Path(args.carpeta).resolve()
    else:
        opciones = [Path(__file__).resolve().parent, Path.cwd()]
        raiz = next((c for c in opciones if any((c / n).is_dir() for n in CARPETAS_ENTRADA)), opciones[0])
    archivos = listar_archivos(raiz)
    if not archivos:
        print("No encontré archivos para anonimizar.")
        return 1

    ruta_clientes = raiz / ARCHIVO_CLIENTES
    if not ruta_clientes.exists():
        print(f"Buscando RUT en {len(archivos)} archivos...")
        filas = detectar_clientes(archivos, raiz, leer_claves_extra(raiz))
        escribir_csv(ruta_clientes, filas, ["rut", "nombre", "tipo", "anonimizar", "veces", "carpetas"])
        sin_nombre = sum(1 for f in filas if not f["nombre"])
        print(f"\nEncontré {len(filas)} RUT y los dejé en {ARCHIVO_CLIENTES}.")
        print("Ábrelo en Excel y revísalo antes de seguir:")
        print("  - Corrige el nombre de cada cliente tal como aparece en las cartolas (nombres y apellidos).")
        if sin_nombre:
            print(f"  - {sin_nombre} RUT no tienen nombre: complétalos si son clientes.")
        print("  - Revisa la columna \"anonimizar\": \"si\" para clientes, \"no\" para bancos, corredoras y emisores.")
        print("    Ya dejé \"no\" en los que parecen instituciones; confírmalo.")
        print("  - Puedes agregar filas con clientes que solo aparecen por nombre (sin RUT).")
        print("Después vuelve a correr este mismo comando.")
        return 0

    nuevos = agregar_ruts_nuevos(ruta_clientes, archivos, raiz, leer_claves_extra(raiz))
    if nuevos:
        print(f"Encontré {len(nuevos)} RUT que no estaban en {ARCHIVO_CLIENTES} y los agregué al final")
        print("(columna \"carpetas\": dónde aparecen). Revísalos y vuelve a correr este comando.")
        return 0

    clientes = leer_clientes(ruta_clientes)
    if not clientes:
        print(f"{ARCHIVO_CLIENTES} no tiene clientes para anonimizar.")
        return 1
    rng = random.Random(args.semilla)
    ruta_mapeo = raiz / ARCHIVO_MAPEO
    asignar_falsos(clientes, leer_mapeo(ruta_mapeo), rng)
    guardar_mapeo(ruta_mapeo, clientes)
    reemplazador = Reemplazador(clientes)
    claves = claves_posibles([c.rut for c in clientes if c.rut], leer_claves_extra(raiz))

    ruta_otros = raiz / ARCHIVO_OTROS
    if not ruta_otros.exists():
        print("Buscando direcciones, comunas, emails, teléfonos y cuentas de los clientes...")
        filas = detectar_otros(archivos, raiz, reemplazador, claves)
        escribir_csv(ruta_otros, filas, COLUMNAS_OTROS)
        por_tipo = Counter(f["tipo"] for f in filas)
        print(f"\nEncontré {len(filas)} datos y los dejé en {ARCHIVO_OTROS}: "
              + ", ".join(f"{n} {tipo}" for tipo, n in sorted(por_tipo.items())))
        print("Revísalo antes de seguir: \"no\" en lo que no sea del cliente, agrega lo que falte")
        print("y los alias de los nombres de archivo (tipo \"alias\"). Después vuelve a correr este comando.")
        return 0
    filas_otros = leer_otros(ruta_otros)
    otros = OtrosDatos(filas_otros, clientes, rng)
    if otros.generados:
        escribir_csv(ruta_otros, [{k: f.get(k, "") for k in COLUMNAS_OTROS} for f in filas_otros], COLUMNAS_OTROS)
    verificador = Verificador(clientes, otros)

    salida = raiz / CARPETA_SALIDA
    if salida.exists():
        shutil.rmtree(salida)
    salida.mkdir()

    reporte, resumen = [], []
    copiados = fallidos = 0
    print(f"Anonimizando {len(archivos)} archivos de {len(clientes)} clientes...")
    for archivo in archivos:
        relativo = archivo.relative_to(raiz)
        partes = [nombre_falso(p, reemplazador, otros) for p in relativo.parts]
        destino = salida.joinpath(*partes)
        n = 2
        while destino.exists():  # dos archivos que quedarían con el mismo nombre
            partes[-1] = f"{Path(partes[-1]).stem} ({n}){Path(partes[-1]).suffix}"
            destino = salida.joinpath(*partes)
            n += 1
        ext = archivo.suffix.lower()
        try:
            if ext in EXT_PDF:
                cambios, problemas, avisos = anonimizar_pdf(archivo, destino, reemplazador, verificador, claves, otros)
            elif ext in EXT_EXCEL:
                cambios, problemas, avisos = anonimizar_excel(archivo, destino, reemplazador, verificador, otros)
            elif ext in EXT_TEXTO:
                cambios, problemas, avisos = anonimizar_texto(archivo, destino, reemplazador, verificador, otros)
            else:
                cambios, problemas, avisos = 0, [f"tipo de archivo {ext or 'sin extensión'} no soportado"], []
        except Exception as e:
            cambios, problemas, avisos = 0, [f"error al procesar: {e}"], []
            if destino.exists():
                destino.unlink()
        problemas_nombre = verificador.revisar_nombre_archivo(Path(*partes))
        if problemas_nombre:
            problemas = problemas + problemas_nombre
            if destino.exists():
                destino.unlink()

        destino_rel = Path(*partes)
        if problemas:
            fallidos += 1
            reporte.append(f"NO COPIADO  {relativo}: {'; '.join(problemas)}")
            resumen.append(f"NO COPIADO  {destino_rel}: {'; '.join(p.split(':')[0] for p in problemas)}")
        else:
            copiados += 1
            linea = f"OK          {relativo} -> {destino_rel} ({cambios} reemplazos)"
            if avisos:
                linea += f"  revisar nombres sueltos: {', '.join(avisos)}"
            reporte.append(linea)
            resumen.append(f"OK          {destino_rel} ({cambios} reemplazos)" + ("  con avisos" if avisos else ""))

    encabezado = [
        f"Archivos revisados: {len(archivos)}",
        f"Copiados a {CARPETA_SALIDA}: {copiados}",
        f"No copiados (requieren revisión): {fallidos}",
        f"Clientes anonimizados: {len(clientes)}",
        "",
    ]
    (raiz / ARCHIVO_REPORTE).write_text("\n".join(encabezado + reporte) + "\n", encoding="utf-8")
    (salida / ARCHIVO_RESUMEN).write_text("\n".join(encabezado + resumen) + "\n", encoding="utf-8")

    print("\n".join(encabezado))
    if fallidos:
        print(f"Los archivos no copiados y el motivo están en {ARCHIVO_REPORTE}.")
    print(f"La carpeta {CARPETA_SALIDA} no tiene nombres, RUT ni otros datos reales de los clientes "
          f"de {ARCHIVO_CLIENTES} y {ARCHIVO_OTROS}.")
    print(f"No compartas {ARCHIVO_CLIENTES}, {ARCHIVO_OTROS}, {ARCHIVO_MAPEO} ni {ARCHIVO_REPORTE}: "
          "tienen datos reales.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
