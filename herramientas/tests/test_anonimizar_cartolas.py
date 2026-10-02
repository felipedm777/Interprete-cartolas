"""Pruebas del anonimizador con cartolas sintéticas (ningún dato real)."""

import csv
import subprocess
import sys
from pathlib import Path

import openpyxl
import openpyxl.comments
import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import anonimizar_cartolas as A  # noqa: E402

SCRIPT = Path(A.__file__)
RUT_PERSONA, RUT_EMPRESA, RUT_BANCO = 12345678, 76123456, 97004000


def rut(cuerpo):
    return f"{cuerpo:,}".replace(",", ".") + "-" + A.digito_verificador(cuerpo)


def pdf_desde_html(ruta, html):
    historia = pymupdf.Story(html=html)
    escritor = pymupdf.DocumentWriter(str(ruta))
    hay_mas = True
    while hay_mas:
        pagina = escritor.begin_page(pymupdf.paper_rect("a4"))
        hay_mas, _ = historia.place(pymupdf.paper_rect("a4") + (36, 36, -36, -36))
        historia.draw(pagina)
        escritor.end_page()
    escritor.close()


@pytest.fixture
def carpeta(tmp_path):
    banco = tmp_path / "Training Data" / "Banco Uno"
    corredora = tmp_path / "Training Data" / "Corredora Dos"
    ejemplo = tmp_path / "Example"
    for c in (banco, corredora, ejemplo):
        c.mkdir(parents=True)

    pdf_desde_html(banco / "Cartola Juan Perez 2026-09.pdf", f"""
        <h2>Banco Uno S.A. RUT {rut(RUT_BANCO)}</h2>
        <p><b>Cliente:</b> JUAN ANDRÉS PÉREZ GONZÁLEZ &nbsp; <b>RUT:</b> {rut(RUT_PERSONA)}</p>
        <p>Señor(a): PÉREZ GONZÁLEZ, JUAN ANDRÉS<br>Av. Apoquindo 1234, Las Condes</p>
        <p>Estimado Sr. Pérez, adjuntamos su cartola del período.</p>
        <table><tr><td>03/09/2026</td><td>Transferencia de Juan Pérez</td><td>250.000</td></tr>
        <tr><td>05/09/2026</td><td>Pago tarjeta</td><td>100.000</td></tr></table>""")

    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((50, 60), "Corredora Dos S.A. Estado de cuenta", fontsize=14, fontname="hebo")
    p.insert_text((50, 90), f"Titular: INVERSIONES LOS ROBLES SPA  R.U.T. {RUT_EMPRESA}-{A.digito_verificador(RUT_EMPRESA)}", fontsize=10)
    p.insert_text((50, 105), f"Apoderado: Juan Andrés Pérez González  RUT {rut(RUT_PERSONA)}", fontsize=10)
    p.insert_text((50, 130), "CFIMRCLP   1.000   12.345   12.345.000", fontsize=10, fontname="cour")
    p.insert_text((30, 700), f"Folio cliente {rut(RUT_PERSONA)}", fontsize=8, rotate=90)
    rotada = doc.new_page()
    rotada.insert_text((50, 60), "Cliente JUAN PEREZ GONZALEZ", fontsize=10)
    rotada.set_rotation(90)
    doc.set_metadata({"author": "JUAN ANDRES PEREZ GONZALEZ"})
    doc.save(corredora / "protegida.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256,
             user_pw=str(RUT_PERSONA), owner_pw="dueno")

    imagen = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 400), False)
    imagen.clear_with(200)
    escaneada = pymupdf.open()
    escaneada.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), pixmap=imagen)
    escaneada.save(corredora / "escaneada.pdf")

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Saldos"
    hoja.append(["RUT", "DV", "Nombre", "Instrumento", "Saldo"])
    hoja.append([RUT_PERSONA, A.digito_verificador(RUT_PERSONA), "Pérez González Juan Andrés", "CFIMRCLP", 12345000])
    hoja.append([RUT_EMPRESA, A.digito_verificador(RUT_EMPRESA), "Inversiones Los Robles SpA", "CLP", 500])
    movs = libro.create_sheet("Movimientos")
    movs.append(["Rut Cliente", "Fecha", "Glosa", "Monto"])
    movs.append([rut(RUT_PERSONA), "2026-09-03", "Abono de JUAN PEREZ", 250000])
    movs["C2"].comment = openpyxl.comments.Comment("Revisado con Juan Pérez González", "Felipe")
    libro.save(ejemplo / "saldos_movimientos.xlsx")
    return tmp_path


def correr(carpeta):
    resultado = subprocess.run(
        [sys.executable, str(SCRIPT), "--carpeta", str(carpeta), "--semilla", "7"],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    return resultado.stdout


def leer_clientes_csv(carpeta):
    with open(carpeta / "clientes.csv", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f, delimiter=";"))


def texto_pdf(ruta):
    doc = pymupdf.open(ruta)
    texto = "\n".join(p.get_text("text") for p in doc) + str(doc.metadata)
    doc.close()
    return A.normalizar(texto)


def test_digito_verificador():
    assert A.digito_verificador(12345678) == "5"
    assert A.digito_verificador(97004000) == "5"
    assert A.digito_verificador(11111111) == "1"
    assert A.digito_verificador(10000013) == "K"


def test_primera_pasada_propone_clientes(carpeta):
    salida = correr(carpeta)
    assert "clientes.csv" in salida
    assert not (carpeta / "anonimizado").exists()
    filas = {f["rut"]: f for f in leer_clientes_csv(carpeta)}
    assert filas[rut(RUT_PERSONA)]["nombre"] == "JUAN ANDRÉS PÉREZ GONZÁLEZ"
    assert filas[rut(RUT_PERSONA)]["anonimizar"] == "si"
    assert filas[rut(RUT_BANCO)]["anonimizar"] == "no"  # el banco no es cliente


def test_segunda_pasada_anonimiza_y_verifica(carpeta):
    correr(carpeta)
    # El PDF con clave se abrió con el RUT del cliente encontrado en otro archivo.
    filas = {f["rut"]: f for f in leer_clientes_csv(carpeta)}
    assert filas[rut(RUT_EMPRESA)]["nombre"] == "INVERSIONES LOS ROBLES SPA"

    correr(carpeta)
    salida = carpeta / "anonimizado"
    pdfs = sorted(salida.rglob("*.pdf"))
    nombres_archivo = [p.name for p in pdfs]
    assert "escaneada.pdf" not in nombres_archivo  # no se puede verificar: no se copia
    assert not any("Perez" in n or "Juan" in n for n in nombres_archivo)
    assert len(pdfs) == 2

    for pdf in pdfs:
        texto = texto_pdf(pdf)
        assert "PEREZ" not in texto and "GONZALEZ" not in texto and "ROBLES" not in texto
        assert "12345678" not in texto.replace(".", "")
        assert "76123456" not in texto.replace(".", "")
    banco = texto_pdf(next(p for p in pdfs if "Banco Uno" in str(p)))
    assert "APOQUINDO 1234" in banco  # el resto del texto queda intacto
    assert rut(RUT_BANCO) in banco  # el RUT del banco no se toca
    assert "TRANSFERENCIA" in banco and "250.000" in banco
    assert not pymupdf.open(next(p for p in pdfs if "protegida" in p.name)).needs_pass

    libro = openpyxl.load_workbook(salida / "Example" / "saldos_movimientos.xlsx")
    fila = [c.value for c in libro["Saldos"][2]]
    assert fila[0] != RUT_PERSONA and fila[1] == A.digito_verificador(fila[0])
    assert "Pérez" not in fila[2] and fila[3] == "CFIMRCLP" and fila[4] == 12345000
    assert libro["Saldos"][3][2].value.startswith("Inversiones Los ") and libro["Saldos"][3][2].value.endswith(" SpA")
    glosa = libro["Movimientos"][2][2]
    assert glosa.value.startswith("Abono de ") and "PEREZ" not in glosa.value
    assert glosa.comment.text.startswith("Revisado con ") and "Pérez" not in glosa.comment.text

    resumen = (salida / "RESUMEN.txt").read_text(encoding="utf-8")
    assert "NO COPIADO" in resumen and "escaneada.pdf" in resumen
    assert "PEREZ" not in A.normalizar(resumen) and "12345678" not in resumen

    # Correr de nuevo usa los mismos datos falsos.
    antes = (carpeta / "mapeo_anonimizacion.csv").read_text(encoding="utf-8-sig")
    correr(carpeta)
    assert (carpeta / "mapeo_anonimizacion.csv").read_text(encoding="utf-8-sig") == antes


def test_reemplazo_en_texto_mantiene_estilo():
    cliente = A.Cliente(RUT_PERSONA, "Juan Andrés Pérez González", "persona")
    A.asignar_falsos([cliente], {}, __import__("random").Random(1))
    r = A.Reemplazador([cliente])
    texto, cambios = r.reemplazar(f"PEREZ GONZALEZ, JUAN ANDRES RUT {rut(RUT_PERSONA)}; juan pérez")
    assert cambios == 3
    assert "PEREZ" not in texto and "juan" not in texto.lower()
    apellidos, nombres = texto.split(" RUT ")[0].split(", ")
    assert apellidos.isupper() and nombres.isupper()
    nuevo_rut = texto.split(" RUT ")[1].split(";")[0]
    cuerpo, dv = nuevo_rut.replace(".", "").split("-")
    assert A.digito_verificador(int(cuerpo)) == dv and nuevo_rut.count(".") == 2
